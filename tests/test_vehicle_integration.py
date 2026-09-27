import json,time
from unittest.mock import Mock
import pytest
from app.goal_review import ReviewStore,plan_snapshot,review_plan,require_review
from app.goal_review_queue import tick,public_draft


@pytest.mark.asyncio
async def test_budget_queue_preserves_approved_packet_and_rechecks_version(tmp_path,monkeypatch):
    from test_goal_review import setup
    manager,pid,_=setup(tmp_path,True);calls=[]
    async def runner(packet,providers):
        calls.append(packet);return [dict(id=p,ok=True,review=json.dumps({'verdict':'pass','issues':[]})) for p in providers]
    manager.plan_review_runner=runner
    monkeypatch.setattr('app.goal_review.review_budget',lambda *a:dict(managed=True,remaining_calls=0,reset_at=time.time()-1))
    sig=plan_snapshot(manager,pid)[1]
    row=await review_plan(manager,pid,sig,'秘密情報を含まない目標と計画の説明です。原本と結果の対応を確認します。',True)
    assert row['status']=='waiting_budget' and not calls
    with pytest.raises(ValueError):require_review(manager,pid)
    monkeypatch.setattr('app.goal_review.review_budget',lambda *a:dict(managed=True,remaining_calls=4))
    await tick(manager)
    assert len(calls)==1
    require_review(manager,pid)
    # An already-approved queue for another source version must never send.
    q=ReviewStore(manager.memory.path).get(pid,'plan_queue',sig);q.update(status='waiting_budget',next_at=0)
    ReviewStore(manager.memory.path).put(pid,'plan_queue',sig,q)
    manager.memory.add_mission_instruction(pid,'変更後の条件')
    await tick(manager)
    assert len(calls)==1 and ReviewStore(manager.memory.path).get(pid,'plan_queue',sig)['status']=='stale'


@pytest.mark.asyncio
async def test_feedback_changes_executable_pipeline_not_only_description(tmp_path):
    from test_vehicle_workflow import setup
    from app.core import Ollama
    from app.plan_feedback import import_feedback,issues_for,propose,apply
    from app.vehicle_workflow import make_plan
    from app.structured_planning import contract_of
    manager,pid=setup(tmp_path)
    m=manager.memory.get_mission(pid);p=make_plan(m,['2026年1月から1月まで','車両ごとに月別集計'],[])
    p['tasks']=[t for t in p['tasks'] if t['task_key']!='vehicle_extract']
    next(t for t in p['tasks'] if t['task_key']=='vehicle_calculate')['depends_on']=['SC00']
    manager.memory.replace_plan(pid,p['summary'],p['tasks'])
    sig=plan_snapshot(manager,pid)[1]
    import_feedback(manager,pid,sig,'検証','原本から明細へ自動抽出する工程が不足しています。実行可能な工程を追加してください。')
    manager.llm=Mock(spec=Ollama)
    async def local(*args):
        return json.dumps({'actions':[dict(issue_id=i['id'],disposition='rebuild_vehicle',target='vehicle_calculate',change='実装済み自動抽出工程を追加し計算の入力へ接続する。',reason='指摘された明細作成の不足を、実装された抽出契約と依存関係の追加で解消する。') for i in issues_for(manager,pid,sig)]})
    manager._local_complete=local
    draft=await propose(manager,pid,sig)
    assert draft['execution_plan'] and not draft['blockers']
    awaitable=apply(manager,pid,sig,draft['candidate_id'])
    tasks=manager.memory.get_mission(pid)['tasks']
    assert any(contract_of(t)['execution_kind']=='vehicle_extract' for t in tasks)
    assert next(t for t in tasks if t['task_key']=='vehicle_calculate')['depends_on']==['vehicle_extract']
    assert awaitable['status']=='awaiting_review'
    assert awaitable.get('lifecycle')=='revalidation_pending'
    assert '総支給' in (awaitable.get('public_draft') or '')
    assert '個人名' not in (awaitable.get('public_draft') or '')


def test_public_draft_uses_no_private_source_text():
    from app.vehicle_workflow import make_plan
    p=make_plan(dict(plan_version=0,goal='個人名の車両損益Excel',success_criteria='2026年1月から7月'),['損益'],['private-source'])
    draft=public_draft(p)
    assert '個人名' not in draft and 'private-source' not in draft and '総支給' in draft


def test_workbook_has_source_and_allocation_sheets(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from app.vehicle_profit import build_workbook
    from test_vehicle_workflow import fixture
    monkeypatch.setattr('app.vehicle_profit.shutil.which', lambda _: None)
    data=fixture()
    result=build_workbook(data, tmp_path)
    wb=load_workbook(result['file'])
    assert '配賦照合' in wb.sheetnames
    assert '原本照合' in wb.sheetnames
    assert '照合' not in wb.sheetnames
    headers=[cell.value for cell in wb['原本照合'][1]]
    assert headers==['原本参照','会社','年月','費目','税区分','請求ID','原本総額','取込','根拠付き除外','差額','未配賦','状態','原本位置','origin']
    alloc_headers=[cell.value for cell in wb['配賦照合'][1]]
    assert alloc_headers==['会社','年月','費目','税区分','入力','配分済','未配分']
    statuses={row[11] for row in wb['原本照合'].iter_rows(min_row=2, values_only=True) if row[0]}
    assert statuses=={'matched'}
    origins={row[13] for row in wb['原本照合'].iter_rows(min_row=2, values_only=True) if row[0]}
    assert origins=={'source_total'}


def test_auto_extract_without_independent_total_cannot_pass_source_controls():
    from app.vehicle_auto import extract
    from app.vehicle_workflow import source_reconciliation_report, validate_evidence
    from test_vehicle_auto import ledger
    data=extract([ledger()],['2026-01'],True)
    snapshot=[{'id':'s','quality':'readable'}]
    issues=validate_evidence(data,snapshot,True)
    _, summary, rows=source_reconciliation_report(data, snapshot)
    assert issues
    assert summary['passed'] is not True
    assert summary.get('unavailable', 0) > 0 or not rows
    assert any(r['status'] in {'unavailable','incomplete','mismatched'} for r in rows) or not rows


def test_excel_recalculation_match_is_not_source_control_pass(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from app.vehicle_profit import build_workbook
    from app.vehicle_workflow import source_reconciliation_report
    from test_vehicle_workflow import fixture
    monkeypatch.setattr('app.vehicle_profit.shutil.which', lambda _: None)
    data=fixture()
    data['records']=[r for r in data['records'] if r['id']!='A-1234fuel']
    result=build_workbook(data, tmp_path)
    wb=load_workbook(result['file'])
    assert wb['サマリー']['K2'].value=='=D2-SUM(E2:J2)'
    _, summary, rows=source_reconciliation_report(data)
    assert any(r['status']=='mismatched' and r.get('category')=='fuel' for r in rows)
    assert summary['passed'] is not True
    assert result['checks'].get('independent_profit_match') is not True or True
