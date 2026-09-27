"""Explicit normalized inputs, exact IDs and Decimal reconciliation. No inferred rates."""
import json
import re
import shutil
import subprocess
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from collections import defaultdict


def number(value):
    if isinstance(value,bool):raise ValueError('金額・比率は数値で指定してください')
    try:n=Decimal(str(value))
    except InvalidOperation:raise ValueError('数値が不正です')
    if not n.is_finite() or abs(n)>Decimal('1e14'):raise ValueError('数値が範囲外です')
    return n


def safe_text(cell,value):
    cell.value=str(value)
    cell.data_type='s'


def calculate(data):
    if data.get('schema')!='vehicle-profit-input-v1':raise ValueError('vehicle-profit-input-v1 の正規化入力が必要です')
    if data.get('rounding')!='yen_half_up':raise ValueError('端数規則 yen_half_up を明示してください')
    if not data.get('basis_note'):raise ValueError('税区分・期間・推計方針の説明が必要です')
    vehicles=data.get('vehicles',[]); months=data.get('months',[]); records=data.get('records',[])
    if not vehicles or len(vehicles)>200 or not months or len(months)>36 or len(records)>30000:raise ValueError('入力件数が範囲外です')
    ids=[v['id'] for v in vehicles]
    if len(set(ids))!=len(ids) or any(not isinstance(i,str) or not i or len(i)>100 for i in ids):raise ValueError('車両IDは一意な文字列が必要です')
    if len(set(months))!=len(months) or any(not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',m) for m in months):raise ValueError('年月が不正です')
    totals=defaultdict(lambda:Decimal(0)); groups=defaultdict(lambda:Decimal(0)); seen=set(); details=[]; warnings=[]
    for r in records:
        rid=r.get('id'); month=r.get('month'); category=r.get('category'); company=r.get('company'); basis=r.get('tax_basis')
        if not isinstance(rid,str) or not rid or rid in seen:raise ValueError('明細IDが重複または未指定です')
        seen.add(rid)
        if month not in months or category not in {'revenue','payroll','insurance','fuel','toll','lease','other'} or not company or basis not in {'exclusive','inclusive','non_taxable','tax_component'}:raise ValueError('月・費目・会社・税区分を確認してください')
        if r.get('quality') not in {'actual','estimated','missing','assumed_zero'}:raise ValueError('実額・推計・欠測の区分が必要です')
        if r['quality']=='missing':
            if any(a.get('vehicle_id') not in ids for a in r.get('allocations',[])):raise ValueError('欠測の対象車両が不明です')
            if r.get('amount') is not None:raise ValueError('欠測金額をゼロで代用できません')
            warnings.append(rid+': 金額欠測');details.append(dict(r,allocated=None,unallocated=None));continue
        amount=number(r.get('amount'))
        if r['quality']=='assumed_zero' and (amount!=0 or data.get('missing_policy')!='zero'):
            raise ValueError('仮定ゼロには不足0円方針と金額0が必要です')
        if amount!=amount.quantize(Decimal('1')):raise ValueError('円単位の整数金額が必要です')
        if r['quality']=='estimated' and not r.get('evidence'):raise ValueError('推計根拠が必要です')
        if r['quality']=='estimated':warnings.append(rid+': 推計')
        if not r.get('source_ref'):raise ValueError('入力原本への参照が必要です')
        allocations=r.get('allocations',[]); allocated=Decimal(0); ratio_total=Decimal(0); target_seen=set()
        for a in allocations:
            if a.get('vehicle_id') not in ids or a['vehicle_id'] in target_seen:raise ValueError('車両IDが不明または配分先重複です。末尾だけでは結合しません')
            target_seen.add(a['vehicle_id']);ratio=number(a.get('ratio'));ratio_total+=ratio
            if not 0<ratio<=1 or ratio_total>1:raise ValueError('配分率は正数、合計1以下が必要です')
            val=(amount*ratio).quantize(Decimal('1'),rounding=ROUND_HALF_UP);allocated+=val
            totals[(a['vehicle_id'],month,category)]+=val
        if abs(allocated)>abs(amount):raise ValueError('丸め後配分額が元金額を超えています。配分を調整してください')
        residual=amount-allocated
        if residual:warnings.append(rid+': 未配分 '+str(residual)+'円')
        groups[(company,month,category,basis,'input')]+=amount
        groups[(company,month,category,basis,'allocated')]+=allocated
        groups[(company,month,category,basis,'unallocated')]+=residual
        details.append(dict(r,allocated=int(allocated),unallocated=int(residual)))
    if not records:raise ValueError('明細がありません')
    # Tax bases remain explicit. Mixed inclusive/exclusive must not become comparable profits.
    if len({r.get('tax_basis') for r in records}&{'inclusive','exclusive'})>1:warnings.append('税込・税抜が混在しています。比較基準を統一してください')
    rows=[]
    for v in vehicles:
        for month in months:
            values={c:int(totals[(v['id'],month,c)]) for c in ['revenue','payroll','insurance','fuel','toll','lease','other']}
            missing={r['category'] for r in records if r['month']==month and r['quality']=='missing'
                     and (not r.get('allocations') or any(a.get('vehicle_id')==v['id'] for a in r['allocations']))}
            if data.get('required_categories'):
                for category in data['required_categories']:
                    if not any(r['month']==month and r['category']==category and r['quality']!='missing'
                               and any(a.get('vehicle_id')==v['id'] for a in r.get('allocations',[])) for r in records):
                        missing.add(category)
                for detail in details:
                    if detail['month']==month and detail.get('unallocated'):
                        missing.add(detail['category'])
                bases={r['tax_basis'] for r in records if r['month']==month and r['quality']!='missing'
                       and any(a.get('vehicle_id')==v['id'] for a in r.get('allocations',[]))} & {'inclusive','exclusive'}
                if len(bases)>1:missing.update(data['required_categories'])
                if missing:warnings.append(v['id']+'/'+month+': 未確定費目 '+','.join(sorted(missing)))
            if not data.get('provisional_numeric'):
                for category in missing:values[category]=None
            profit=None if missing and not data.get('provisional_numeric') else values['revenue']-sum(values[c] for c in values if c!='revenue')
            rows.append(dict(vehicle=v['id'],label=v.get('label',v['id']),month=month,**values,profit=profit,quality='暫定・要確認' if warnings else ('仮定ゼロあり' if any(r['quality']=='assumed_zero' and r['month']==month and any(a['vehicle_id']==v['id'] for a in r.get('allocations',[])) for r in records) else '照合対象')))
    reconciliation=[]
    for key in sorted({k[:4] for k in groups}):
        a=groups[key+('input',)];b=groups[key+('allocated',)];c=groups[key+('unallocated',)]
        if a!=b+c:raise ValueError('入力と配分の合計が不一致です')
        reconciliation.append(dict(company=key[0],month=key[1],category=key[2],tax_basis=key[3],input=int(a),allocated=int(b),unallocated=int(c)))
    return dict(rows=rows,details=details,reconciliation=reconciliation,warnings=warnings)


def _source_reconciliation_rows(data, rows=None):
    if rows is not None:
        return rows
    from app.vehicle_reconciliation import reconcile_sources
    records=[r for r in data.get('records',[]) if r.get('quality') not in {'missing','assumed_zero'} and not str(r.get('source_ref') or '').startswith('policy:')]
    read_issues=[i for i in (data.get('auto_extraction') or {}).get('issues',[]) if i.get('kind')=='read' and i.get('status')!='resolved']
    return reconcile_sources(records, list(data.get('source_controls') or []), list(data.get('source_dispositions') or []), list(data.get('excluded_records') or []), read_issues)


def _control_locator(data, row):
    for control in data.get('source_controls') or []:
        if all((control.get(k) or None)==(row.get(k) or None) for k in ('source_ref','company','month','category','tax_basis','invoice_id')):
            return control.get('source_locator') or ''
    return ''


def build_workbook(data,directory,cancelled=lambda:False,source_reconciliation=None):
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    result=calculate(data); directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    wb=Workbook();summary=wb.active;summary.title='サマリー'
    summary.append(['車両ID','車両名','年月','売上','給与','会社負担保険','燃料','高速','リース（税込）','その他','利益','状態'])
    for n,r in enumerate(result['rows'],2):
        summary.append([r['vehicle'],r['label'],r['month'],r['revenue'],r['payroll'],r['insurance'],r['fuel'],r['toll'],r['lease'],r['other'],(f'=D{n}-SUM(E{n}:J{n})' if r['profit'] is not None else None),r['quality']])
        for col in (1,2,3,12):safe_text(summary.cell(n,col),summary.cell(n,col).value)
    for idx,v in enumerate(data['vehicles'],1):
        sheet=wb.create_sheet(f'車両_{idx:03d}');sheet.append([v['id'],v.get('label',v['id'])]);safe_text(sheet['A1'],v['id']);safe_text(sheet['B1'],v.get('label',v['id']))
        sheet.append(['年月','売上','給与','会社負担保険','燃料','高速','リース（税込）','その他','利益','状態'])
        for n,r in enumerate(result['rows'],2):
            if r['vehicle']==v['id']:
                sheet.append([None if summary[f'{col}{n}'].value is None else f"='サマリー'!{col}{n}" for col in 'CDEFGHIJKL'])
    detail=wb.create_sheet('入力明細');detail.append(['明細ID','会社','年月','費目','税区分','区分','入力金額','配分済','未配分','原本参照','推計根拠'])
    for r in result['details']:
        detail.append([r['id'],r['company'],r['month'],r['category'],r['tax_basis'],r['quality'],r.get('amount'),r['allocated'],r['unallocated'],r.get('source_ref',''),r.get('evidence','')])
        for col in (1,2,3,4,5,6,10,11):safe_text(detail.cell(detail.max_row,col),detail.cell(detail.max_row,col).value)
    reconcile=wb.create_sheet('配賦照合');reconcile.append(['会社','年月','費目','税区分','入力','配分済','未配分'])
    for r in result['reconciliation']:
        reconcile.append(list(r.values()))
        for col in range(1,5):safe_text(reconcile.cell(reconcile.max_row,col),reconcile.cell(reconcile.max_row,col).value)
    source_sheet=wb.create_sheet('原本照合');source_sheet.append(['原本参照','会社','年月','費目','税区分','請求ID','原本総額','取込','根拠付き除外','差額','未配賦','状態','原本位置','origin'])
    for r in _source_reconciliation_rows(data, source_reconciliation):
        source_sheet.append([
            r.get('source_ref') or '', r.get('company') or '', r.get('month') or '', r.get('category') or '',
            r.get('tax_basis') or '', r.get('invoice_id') or '', r.get('source_total'), r.get('included_total'),
            r.get('excluded_total'), r.get('difference'), r.get('unallocated_total'), r.get('status') or '',
            _control_locator(data, r), r.get('origin') or '',
        ])
        for col in (1,2,3,4,5,6,12,13,14):
            safe_text(source_sheet.cell(source_sheet.max_row,col),source_sheet.cell(source_sheet.max_row,col).value)
    audit=wb.create_sheet('原本対応');audit.append(['明細ID','原本参照','ページ・セル','根拠','区分'])
    for r in result['details']:
        audit.append([r['id'],r.get('source_ref',''),r.get('source_locator',''),r.get('evidence',''),r['quality']])
        for c in audit[audit.max_row]:safe_text(c,c.value)
    exceptions=wb.create_sheet('不足と未配賦');exceptions.append(['対象','年月','費目','金額','状態・理由'])
    for r in result['details']:
        if r['quality']=='assumed_zero' or r.get('unallocated'):
            exceptions.append([r['id'],r['month'],r['category'],r.get('unallocated') or 0,'仮定ゼロ' if r['quality']=='assumed_zero' else '未配賦'])
            for col in (1,2,3,5):safe_text(exceptions.cell(exceptions.max_row,col),exceptions.cell(exceptions.max_row,col).value)
    for r in data.get('excluded_records',[]):
        exceptions.append([r['id'],r['month'],r['category'],r['amount'],'対象外: '+r.get('exclusion_reason','')])
        for col in (1,2,3,5):safe_text(exceptions.cell(exceptions.max_row,col),exceptions.cell(exceptions.max_row,col).value)
    for issue in data.get('auto_extraction',{}).get('issues',[]):
        exceptions.append([issue['source_ref'],'','',issue.get('amount'),issue['message']])
        for col in (1,2,3,5):safe_text(exceptions.cell(exceptions.max_row,col),exceptions.cell(exceptions.max_row,col).value)
    notes=wb.create_sheet('説明');safe_text(notes['A1'],data['basis_note']);notes.append(['原本から抽出した入力または確認済み入力を計算。制度料率を推測しません。仮定ゼロ・未配賦・税基準混在を確認してください。'])
    for warning in result['warnings']:notes.append([warning]);safe_text(notes.cell(notes.max_row,1),warning)
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='214E65')
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=22
        for row in sheet:
            for cell in row:
                if cell.data_type=='n':cell.number_format='#,##0'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
    raw=directory/'profit.xlsx';wb.save(raw)
    if cancelled():return dict(status='cancelled')
    exe=shutil.which('libreoffice') or shutil.which('soffice')
    if not exe:return dict(status='needs_review',warnings=result['warnings']+['再計算機能なし'],file=str(raw),checks={'recalculation':False})
    out=directory/'recalculated';out.mkdir(exist_ok=True)
    command=[exe,'-env:UserInstallation='+ (directory/'lo_profile').resolve().as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),str(raw)]
    proc=subprocess.run(command,capture_output=True,text=True,timeout=120)
    (directory/'recalculation.log').write_text(proc.stdout+'\n'+proc.stderr,encoding='utf-8')
    recalculated=out/'profit.xlsx'
    if proc.returncode or not recalculated.exists():raise ValueError('LibreOffice再計算に失敗しました')
    actual=load_workbook(recalculated,data_only=True)
    errors=[f'{s.title}!{c.coordinate}' for s in actual for row in s for c in row if c.data_type=='e']
    if errors:raise ValueError('数式エラー: '+','.join(errors[:20]))
    for n,r in enumerate(result['rows'],2):
        if actual['サマリー'].cell(n,11).value!=r['profit']:raise ValueError('独立計算値と利益が一致しません')
    for idx,v in enumerate(data['vehicles'],1):
        vr=[r for r in result['rows'] if r['vehicle']==v['id']]
        for n,r in enumerate(vr,3):
            if actual[f'車両_{idx:03d}'].cell(n,9).value!=r['profit']:raise ValueError('車両シートとサマリーが不一致です')
    return dict(status='needs_review' if result['warnings'] else 'completed',warnings=result['warnings'],file=str(recalculated),
        checks={'recalculation':True,'formula_errors':0,'independent_profit_match':True,'vehicle_summary_match':True,'input_allocation_reconciled':True},expected=result['rows'])
