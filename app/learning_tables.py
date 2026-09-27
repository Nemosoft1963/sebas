"""Finite table operations used by learned procedures; never evaluate imported code."""
import ast
import csv
import hashlib
import io
import json
import operator
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

MAX_ROWS=20000
MAX_NODES=40
OPS={'union','filter','derive','join','aggregate','assert','vehicle_profit','export'}
FIELDS=['id','vehicle_id','label','employee_id','company','month','category','tax_basis','quality','amount','ratio','source_ref']
ALIASES={'vehicle_id':['車両ID','車番','車両番号','vehicle','vehicle_id'],'label':['車両名','名称','label'],'employee_id':['社員番号','従業員ID','employee_id'],'company':['会社','会社名','company'],'month':['年月','対象月','月','month'],'amount':['金額','売上','給与総額','合計','amount'],'ratio':['配分率','ratio'],'id':['明細ID','ID','id'],'category':['費目','category'],'tax_basis':['税区分','tax_basis'],'quality':['実額推計区分','quality']}


def fingerprint(value):return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest()

def number(value):
    if isinstance(value,bool) or value is None:raise ValueError('欠測または不正な数値です')
    text=str(value).strip().replace(',','').replace('￥','').replace('¥','')
    if text.startswith('(') and text.endswith(')'):text='-'+text[1:-1]
    try:n=Decimal(text)
    except InvalidOperation:raise ValueError('数値に変換できません: '+text[:40])
    if not n.is_finite() or abs(n)>Decimal('1e14'):raise ValueError('数値が範囲外です')
    return n


def decode(raw):
    for encoding in ('utf-8-sig','cp932','utf-16'):
        try:return raw.decode(encoding)
        except UnicodeError:pass
    raise ValueError('文字コードを読み取れません')


def inspect_table(raw,suffix,sheet=None,header_row=1):
    if type(header_row) is not int or not 1<=header_row<=100:raise ValueError('見出し行は1〜100です')
    sheets=[]
    if suffix=='.xlsx':
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist())>100*1024*1024:raise ValueError('展開サイズが上限を超えます')
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        sheets=wb.sheetnames
        if sheet is None:sheet=sheets[0]
        if sheet not in sheets:raise ValueError('指定シートがありません')
        ws=wb[sheet]
        if ws.max_row>MAX_ROWS+100 or ws.max_column>100:raise ValueError('表の上限を超えます')
        matrix=[list(row) for row in ws.iter_rows(values_only=True)];wb.close()
    elif suffix in {'.csv','.tsv'}:
        text=decode(raw)
        try:delimiter=csv.Sniffer().sniff(text[:8192],delimiters=',\t;').delimiter
        except csv.Error:delimiter='\t' if suffix=='.tsv' else ','
        matrix=list(csv.reader(io.StringIO(text),delimiter=delimiter))
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if len(reader.pages)>100:raise ValueError('PDFは100ページ以下にしてください')
        text='\n'.join(page.extract_text() or '' for page in reader.pages)
        if not text.strip():return {'status':'needs_capability','reason':'画像PDFのためOCR能力が必要です','sheets':[],'headers':[],'rows':[]}
        # Only delimiter-preserving tables qualify; ordinary prose is never guessed into money.
        lines=[line for line in text.splitlines() if '\t' in line or '|' in line]
        if len(lines)<2:return {'status':'needs_capability','reason':'表の列境界を確定できません。CSV/Excelか形式別抽出器が必要です','sheets':[],'headers':[],'rows':[]}
        matrix=[re.split(r'\t|\|',line.strip('|')) for line in lines]
    else:raise ValueError('表はCSV/TSV/XLSX/PDFを指定してください')
    if len(matrix)>MAX_ROWS+100 or not matrix or header_row>len(matrix):raise ValueError('表の行数が不正です')
    headers=[str(c).strip() if c is not None else '' for c in matrix[header_row-1]]
    while headers and not headers[-1]:headers.pop()
    if not headers or len(headers)>100 or any(not h for h in headers) or len(set(headers))!=len(headers):raise ValueError('空欄・重複のない列名を持つ見出し行を指定してください')
    rows=[]
    for rownum,row in enumerate(matrix[header_row:],header_row+1):
        if not any(v is not None and str(v).strip() for v in row):continue
        if any(v is not None and str(v).strip() for v in row[len(headers):]):raise ValueError('見出しより多い列があります')
        rows.append({'values':{h:(str(row[i]) if i<len(row) and row[i] is not None else None) for i,h in enumerate(headers)},'row':rownum})
    if len(rows)>MAX_ROWS:raise ValueError('入力行数上限を超えます')
    candidates={k:[h for h in headers if h.lower() in [a.lower() for a in aliases]] for k,aliases in ALIASES.items()}
    return {'status':'ready_for_mapping','headers':headers,'rows':rows,'sheets':sheets,'sheet':sheet,'header_row':header_row,'candidates':candidates,'row_count':len(rows)}


def normalize(table,mapping,asset_hash):
    if table['status']!='ready_for_mapping':raise ValueError(table.get('reason','表を読めません'))
    if not mapping.get('confirmed') or not mapping.get('reason','').strip():raise ValueError('列対応と根拠を確認してください')
    columns=mapping.get('columns',{});constants=mapping.get('constants',{})
    if not columns or not all(isinstance(k,str) and re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,49}',k) for k in columns):raise ValueError('対応先は英数字の項目名にしてください')
    if set(columns)&set(constants):raise ValueError('列と固定値が同じ項目を指定しています')
    if not set(columns.values())<=set(table['headers']):raise ValueError('対応先の列が入力にありません')
    result=[]
    for source in table['rows']:
        row={k:source['values'][v] for k,v in columns.items()};row.update(constants)
        row['_source']=asset_hash+':'+str(table.get('sheet') or '')+':'+str(source['row'])
        if 'month' in row and row['month']:
            m=re.fullmatch(r'(\d{4})[年/.-](\d{1,2})(?:月|[-/].*|\s.*)?',str(row['month']))
            if not m or not 1<=int(m[2])<=12:raise ValueError('年月を解釈できません: '+str(row['month']))
            row['month']=f'{m[1]}-{int(m[2]):02d}'
        for field in mapping.get('numeric',[]):
            if row.get(field) not in (None,''):
                val=number(row[field]);factor=number(mapping.get('scales',{}).get(field,1))
                row[field]=str(val*factor)
            else:row[field]=None
        result.append(row)
    return result


def expression(text,row):
    if not isinstance(text,str) or len(text)>300:raise ValueError('式が長すぎます')
    tree=ast.parse(text,mode='eval')
    if sum(1 for _ in ast.walk(tree))>50:raise ValueError('式が複雑すぎます')
    operations={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}
    def visit(node):
        if isinstance(node,ast.Expression):return visit(node.body)
        if isinstance(node,ast.Name):return number(row.get(node.id))
        if isinstance(node,ast.Constant) and type(node.value) in {int,float}:return number(node.value)
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.USub,ast.UAdd)):return -visit(node.operand) if isinstance(node.op,ast.USub) else visit(node.operand)
        if isinstance(node,ast.BinOp) and type(node.op) in operations:
            try:value=operations[type(node.op)](visit(node.left),visit(node.right))
            except (ArithmeticError,InvalidOperation):raise ValueError('ゼロ除算または演算エラー')
            return number(value)
        raise ValueError('式には項目名と数値の四則演算だけを使用できます')
    return str(visit(tree))


def validate_recipe(recipe):
    nodes=recipe.get('nodes',[])
    if not nodes or len(nodes)>MAX_NODES:raise ValueError('工程は1〜40件です')
    known=set(recipe.get('inputs',[]));ids=set()
    if not known or not all(re.fullmatch(r'[a-z][a-z0-9_]{0,49}',k) for k in known):raise ValueError('入力名が不正です')
    for node in nodes:
        nid=node.get('id','');op=node.get('op');inputs=node.get('inputs',[])
        if not re.fullmatch(r'[a-z][a-z0-9_]{0,49}',nid) or nid in known or nid in ids:raise ValueError('工程IDが重複または不正です')
        if op not in OPS:raise ValueError('未対応能力: '+str(op))
        if not inputs or not all(i in known for i in inputs):raise ValueError('工程入力は既存入力または先行工程に限定します')
        if op=='join' and len(inputs)!=2:raise ValueError('結合入力は2つです')
        if op in {'derive','aggregate','assert','filter','export'} and len(inputs)!=1:raise ValueError('この工程の入力は1つです')
        if not isinstance(node.get('params',{}),dict):raise ValueError('工程条件が不正です')
        ids.add(nid);known.add(nid)
    if nodes[-1]['op'] not in {'export','vehicle_profit'}:raise ValueError('最後の工程は表出力または車両損益出力にしてください')
    return recipe


def vehicle_input(tables,params):
    vehicles=tables.get(params.get('vehicles','vehicles'),[])
    if not vehicles:raise ValueError('車両マスタが必要です')
    ids=[v.get('vehicle_id') for v in vehicles]
    if any(not i for i in ids) or len(set(ids))!=len(ids):raise ValueError('車両マスタのIDが重複または欠落しています')
    rows=[]
    for key in params.get('records',[]):rows.extend(tables[key])
    if not rows:raise ValueError('金額明細がありません')
    allocation_rows=tables.get(params.get('allocations','allocations'),[])
    allocs=defaultdict(list)
    for a in allocation_rows:
        key=(a.get('employee_id'),a.get('company'),a.get('month'))
        if not all(key):raise ValueError('配車対応に従業員・会社・月が必要です')
        allocs[key].append({'vehicle_id':a.get('vehicle_id'),'ratio':str(number(a.get('ratio')))})
    records=[];seen=set();months=set()
    for row in rows:
        source=row['_source']
        if source in seen:raise ValueError('同じ原本明細を重複して取り込んでいます')
        seen.add(source);months.add(row.get('month'))
        allocations=[]
        if row.get('vehicle_id'):allocations=[{'vehicle_id':row['vehicle_id'],'ratio':'1'}]
        elif row.get('employee_id'):allocations=allocs.get((row['employee_id'],row.get('company'),row.get('month')),[])
        records.append({'id':row.get('id') or source,'month':row.get('month'),'company':row.get('company'),'category':row.get('category'),'tax_basis':row.get('tax_basis'),'quality':row.get('quality'), 'amount':row.get('amount'),'source_ref':source,'evidence':row.get('evidence',''),'allocations':allocations})
    return {'schema':'vehicle-profit-input-v1','rounding':params.get('rounding'),'basis_note':params.get('basis_note',''), 'vehicles':[{'id':v['vehicle_id'],'label':v.get('label',v['vehicle_id'])} for v in vehicles],'months':sorted(months,key=str),'records':records}


def execute_recipe(recipe,inputs,directory,cancelled=lambda:False,checkpoint=lambda x:None,stop_after=None):
    validate_recipe(recipe)
    if set(recipe['inputs'])-set(inputs):raise ValueError('手順の入力資料が不足しています')
    tables={k:[dict(r) for r in v] for k,v in inputs.items()};trace=[];directory=Path(directory);directory.mkdir(parents=True,exist_ok=True);result=None
    for node in recipe['nodes']:
        if cancelled():return {'status':'cancelled','trace':trace}
        op=node['op'];p=node.get('params',{});source=[tables[i] for i in node['inputs']];rows=[dict(r) for r in source[0]]
        if node.get('when_empty')=='skip' and not rows:out=[]
        elif op=='union':out=[dict(r) for part in source for r in part]
        elif op=='filter':
            if p.get('test') not in {'eq','ne','present','missing'}:raise ValueError('絞込条件が未対応です')
            def matches(r):
                value=r.get(p['field']);test=p['test']
                return value==p.get('value') if test=='eq' else value!=p.get('value') if test=='ne' else value not in (None,'') if test=='present' else value in (None,'')
            out=[r for r in rows if matches(r)]
        elif op=='derive':
            if not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,49}',p.get('field','')):raise ValueError('計算項目名が不正です')
            out=[]
            for r in rows:r[p['field']]=expression(p['expression'],r);out.append(r)
        elif op=='join':
            keys=p.get('keys',[])
            if not keys:raise ValueError('結合キーが必要です')
            index={}
            for r in source[1]:
                key=tuple(r.get(k) for k in keys)
                if any(v in (None,'') for v in key) or key in index:raise ValueError('結合先キーが空欄または重複しています')
                index[key]=r
            out=[]
            for r in rows:
                key=tuple(r.get(k) for k in keys);target=index.get(key)
                if target is None and p.get('missing','error')=='error':raise ValueError('結合先が見つかりません')
                merged=dict(r)
                for k in p.get('fields',[]):
                    if k in merged:raise ValueError('結合で既存列を上書きできません')
                    merged[k]=target.get(k) if target else None
                out.append(merged)
        elif op=='aggregate':
            keys=p.get('keys',[]);values=p.get('values',[]);groups={}
            if not values:raise ValueError('集計する数値項目が必要です')
            for r in rows:
                if not set(keys+values)<=set(r):raise ValueError('集計項目がありません')
                key=tuple(r[k] for k in keys)
                if any(v in (None,'') for v in key):raise ValueError('集計キーが欠測しています')
                group=groups.setdefault(key,{k:v for k,v in zip(keys,key)})
                for field in values:group[field]=str(number(group.get(field,0))+number(r[field]))
            out=list(groups.values())
        elif op=='assert':
            kind=p.get('kind');field=p.get('field')
            if kind=='unique':
                values=[r.get(field) for r in rows]
                if any(v in (None,'') for v in values) or len(set(values))!=len(values):raise ValueError('一意性検査に不合格です')
            elif kind=='nonempty':
                if not rows or any(r.get(field) in (None,'') for r in rows):raise ValueError('必須項目の検査に不合格です')
            elif kind=='sum':
                if sum((number(r.get(field)) for r in rows),Decimal(0))!=number(p.get('expected')):raise ValueError('合計が期待値と不一致です')
            else:raise ValueError('検証条件が未対応です')
            out=rows
        elif op=='vehicle_profit':
            from app.vehicle_profit import build_workbook
            data=vehicle_input({k:tables[k] for k in node['inputs']},p)
            result=build_workbook(data,directory,cancelled);out=result.get('expected',[])
        elif op=='export':
            from openpyxl import Workbook
            from app.vehicle_profit import safe_text
            wb=Workbook();ws=wb.active;ws.title='結果';fields=p.get('fields') or list(dict.fromkeys(k for r in rows for k in r if not k.startswith('_')))
            if not fields or len(fields)>100:raise ValueError('出力列がありません')
            ws.append(fields)
            for c in ws[1]:safe_text(c,c.value)
            for r in rows:
                ws.append([r.get(k) for k in fields])
                for c in ws[ws.max_row]:
                    if isinstance(c.value,str):safe_text(c,c.value)
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=22
            file=directory/'result.xlsx';wb.save(file)
            from openpyxl import load_workbook
            saved=load_workbook(file,data_only=True);matrix=list(saved.active.values);saved.close()
            out=[dict(zip(matrix[0],row)) for row in matrix[1:]]
            result={'status':'completed','file':str(file),'expected':out,'checks':{'declared_operations_completed':True}}
        else:raise ValueError('未対応の操作です')
        if len(out)>MAX_ROWS:raise ValueError('中間結果の行数上限を超えました')
        tables[node['id']]=out
        proof={'step':node['id'],'op':op,'rows':len(out),'hash':fingerprint(out),'sample':out[:10]}
        trace.append(proof);checkpoint(proof)
        (directory/(node['id']+'.json')).write_text(json.dumps(out,ensure_ascii=False,default=str),encoding='utf-8')
        if stop_after==node['id']:return {'status':'paused','trace':trace,'intermediate':out[:100]}
    if result is None:raise ValueError('成果物を生成できませんでした')
    result['trace']=trace;return result


def compare_rows(actual,expected,keys,fields):
    if not keys or not fields or not expected:raise ValueError('比較キー・比較項目・期待結果が必要です')
    def indexed(rows):
        out={}
        for r in rows:
            key=tuple(str(r.get(k,'')) for k in keys)
            if any(not v for v in key) or key in out:raise ValueError('比較キーが空欄または重複しています')
            out[key]=r
        return out
    if any(not set(keys+fields)<=set(row) for row in actual+expected):raise ValueError('比較する項目が成果物または期待結果にありません')
    left=indexed(actual);right=indexed(expected);diff=[]
    for key in sorted(left.keys()|right.keys()):
        if key not in left or key not in right:diff.append({'key':key,'reason':'missing' if key not in left else 'extra'});continue
        for field in fields:
            a=left[key].get(field);b=right[key].get(field)
            if a in (None,'') or b in (None,''):equal=a in (None,'') and b in (None,'')
            else:
                try:equal=number(a)==number(b)
                except ValueError:equal=str(a)==str(b)
            if not equal:diff.append({'key':key,'field':field,'actual':a,'expected':b})
    return {'passed':not diff,'difference_count':len(diff),'differences':diff[:200],'compared_rows':len(right)}
