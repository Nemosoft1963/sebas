"""Positive per-criterion checks and structured vehicle goal failures.

PASS requires a check_id and evidence. Missing materials or unimplemented
checks are UNTESTABLE. Empty failure lists never imply PASS.
"""
from __future__ import annotations

from itertools import product

from app.goal_contract import preview
from app.structured_planning import contract_of
from app.vehicle_workflow import (
    classify_unresolved,
    goal_failure_records,
    load_input,
    read_json,
    resolve,
    source_reconciliation_report,
    sources,
)


REQUIRED_EXPENSE_CATEGORIES = ("revenue", "fuel", "lease", "payroll", "insurance", "other")
MAJOR_TRACE_CATEGORIES = ("revenue", "payroll", "insurance", "fuel", "lease")
PAYROLL_INSURANCE = ("payroll", "insurance")


def _row(criterion_id, status, *, reason_code="", evidence_path="", message="",
         check_id="", evidence_summary=""):
    return {
        "criterion_id": criterion_id,
        "status": status,
        "reason_code": reason_code,
        "evidence_path": evidence_path,
        "message": message,
        "check_id": check_id,
        "evidence_summary": evidence_summary,
    }


def _pass(criterion_id, check_id, evidence_path, evidence_summary, message=""):
    return _row(
        criterion_id, "PASS", check_id=check_id,
        evidence_path=evidence_path, evidence_summary=evidence_summary, message=message,
    )


def _fail(criterion_id, check_id, code, message, evidence_path="", evidence_summary=""):
    return _row(
        criterion_id, "FAIL", reason_code=code, check_id=check_id,
        evidence_path=evidence_path, evidence_summary=evidence_summary, message=message,
    )


def _untestable(criterion_id, check_id, message, evidence_path="", code="UNTESTABLE"):
    return _row(
        criterion_id, "UNTESTABLE", reason_code=code, check_id=check_id,
        evidence_path=evidence_path, message=message, evidence_summary=message,
    )


def _is_gap(record):
    if record.get("quality") in {"missing", "assumed_zero"}:
        return True
    return str(record.get("source_ref") or "").startswith("policy:")


def _actual_records(data, category=None):
    records = []
    for record in (data or {}).get("records") or []:
        if _is_gap(record):
            continue
        if category is not None and record.get("category") != category:
            continue
        records.append(record)
    return records


def _vehicle_ids(data):
    return [v.get("id") for v in (data or {}).get("vehicles") or [] if v.get("id")]


def _months_of(data, contract):
    months = list((data or {}).get("months") or [])
    if months:
        return months
    return list(((contract or {}).get("scope") or {}).get("period") or [])


def _allocation_rules(envelope, data):
    rules = []
    for key in ("assignments", "allocation_rules"):
        rules.extend(list((envelope or {}).get(key) or []))
        rules.extend(list((data or {}).get(key) or []))
    return [r for r in rules if isinstance(r, dict)]


def _approved_rules(rules):
    approved = []
    for rule in rules:
        status = rule.get("status")
        if status in {None, "active", "approved"}:
            approved.append(rule)
    return approved


def _record_allocated(record):
    allocations = list(record.get("allocations") or [])
    if not allocations:
        return False
    try:
        total = sum(float(item.get("ratio") or 0) for item in allocations)
    except (TypeError, ValueError):
        return False
    return total > 0 and all(item.get("vehicle_id") for item in allocations)


def _calc_artifacts(manager, project_id, mission):
    calc_path = ""
    workbook_path = ""
    result = None
    task = next(
        (t for t in (mission.get("tasks") or [])
         if (contract_of(t) or {}).get("execution_kind") == "vehicle_calculate"),
        None,
    )
    if not task:
        return None, "", "", None
    outputs = list((contract_of(task) or {}).get("outputs") or [])
    json_out = next((o.get("path") for o in outputs if str(o.get("path") or "").endswith(".json")), "")
    xlsx_out = next((o.get("path") for o in outputs if str(o.get("path") or "").endswith(".xlsx")), "")
    calc_file = None
    workbook = None
    if json_out:
        try:
            calc_file = resolve(manager, project_id, json_out)
            calc_path = json_out
            if calc_file.exists():
                result = read_json(calc_file)
        except (OSError, ValueError, TypeError, KeyError):
            result = None
    if xlsx_out:
        try:
            workbook = resolve(manager, project_id, xlsx_out)
            workbook_path = xlsx_out
        except (OSError, ValueError, TypeError, KeyError):
            workbook = None
    return result, calc_path, workbook_path, workbook


def collect_context(manager, project_id, contract=None):
    mission = manager.memory.get_mission(project_id)
    contract = contract or preview(manager, project_id) or {}
    envelope = None
    try:
        envelope = load_input(manager, project_id)
    except Exception:
        envelope = None
    data = (envelope or {}).get("data") if envelope else None
    calc, calc_path, workbook_path, workbook = _calc_artifacts(manager, project_id, mission)
    recon_summary = None
    recon_rows = []
    if data is not None:
        try:
            _, recon_summary, recon_rows = source_reconciliation_report(data, sources(manager, project_id))
        except Exception:
            recon_summary = None
            recon_rows = []
    unresolved = None
    if data is not None:
        try:
            unresolved = classify_unresolved(data)
        except Exception:
            unresolved = None
    return {
        "manager": manager,
        "project_id": project_id,
        "mission": mission,
        "contract": contract,
        "envelope": envelope,
        "data": data,
        "calc": calc,
        "calc_path": calc_path,
        "workbook_path": workbook_path,
        "workbook": workbook,
        "recon_summary": recon_summary,
        "recon_rows": recon_rows,
        "unresolved": unresolved,
        "input_path": "vehicle_profit/input.json",
    }


def check_c01_scope(ctx):
    check_id = "check_scope_determined"
    data = ctx["data"]
    contract = ctx["contract"]
    period = _months_of(data, contract)
    vehicles = _vehicle_ids(data)
    evidence = ctx["input_path"] if ctx["envelope"] else ""
    if data is None:
        return _untestable("C01", check_id, "車両入力が無いため対象会社・期間・車両を判定できません")
    if not period:
        return _fail("C01", check_id, "CHECK_MISSING", "対象期間が空です", evidence)
    if not vehicles:
        return _fail("C01", check_id, "CHECK_MISSING", "対象車両集合が空です", evidence)
    companies = {v.get("company") for v in data.get("vehicles") or [] if v.get("company")}
    companies.update(r.get("company") for r in data.get("records") or [] if r.get("company"))
    summary = f"period={period} vehicles={vehicles} companies={sorted(companies)}"
    return _pass("C01", check_id, evidence, summary)


def check_c02_coverage(ctx):
    check_id = "check_vehicle_month_coverage"
    data = ctx["data"]
    if data is None:
        return _untestable("C02", check_id, "車両入力が無いため全車両×全月の行を判定できません")
    vehicles = _vehicle_ids(data)
    months = _months_of(data, ctx["contract"])
    if not vehicles or not months:
        return _fail("C02", check_id, "CHECK_MISSING", "対象車両または対象月が空です", ctx["input_path"])
    present = set()
    for record in _actual_records(data):
        month = record.get("month")
        for item in record.get("allocations") or []:
            vid = item.get("vehicle_id")
            if vid and month:
                present.add((vid, month))
    missing = [f"{vid}/{month}" for vid, month in product(vehicles, months) if (vid, month) not in present]
    summary = f"required={len(vehicles) * len(months)} present={len(present)}"
    if missing:
        return _fail(
            "C02", check_id, "CHECK_MISSING",
            "全車両×全対象月の行が不足: " + ",".join(missing[:12]),
            ctx["input_path"], summary,
        )
    return _pass("C02", check_id, ctx["input_path"], summary)


def check_c03_categories(ctx):
    check_id = "check_category_definitions"
    data = ctx["data"]
    if data is None:
        return _untestable("C03", check_id, "車両入力が無いため費目定義を判定できません")
    defined = set(data.get("required_categories") or [])
    defined.update(r.get("category") for r in data.get("records") or [] if r.get("category"))
    missing = [c for c in REQUIRED_EXPENSE_CATEGORIES if c not in defined]
    summary = f"defined={sorted(defined)}"
    if missing:
        return _fail(
            "C03", check_id, "CHECK_MISSING",
            "費目が未定義: " + ",".join(missing), ctx["input_path"], summary,
        )
    return _pass("C03", check_id, ctx["input_path"], summary)


def check_c04_allocation_rules(ctx):
    check_id = "check_payroll_insurance_allocation_rules"
    data = ctx["data"]
    envelope = ctx["envelope"]
    if data is None:
        return _untestable("C04", check_id, "車両入力が無いため給与・社保の配賦を判定できません")
    targets = [r for r in _actual_records(data) if r.get("category") in PAYROLL_INSURANCE]
    if not targets:
        return _untestable("C04", check_id, "給与・社保の明細が無いため配賦ルール検査を判定できません")
    unallocated = [r.get("id") for r in targets if not _record_allocated(r)]
    rules = _approved_rules(_allocation_rules(envelope, data))
    summary = f"payroll_insurance={len(targets)} approved_rules={len(rules)} unallocated={len(unallocated)}"
    if unallocated:
        return _fail(
            "C04", check_id, "ALLOCATION_UNRESOLVED",
            "未配賦の給与・社保が残っています: " + ",".join(str(x) for x in unallocated[:12]),
            ctx["input_path"], summary,
        )
    if not rules:
        return _untestable(
            "C04", check_id,
            "承認済み配賦ルール(allocation_rules/assignments)が無いため給与・社保の規則集計を判定できません",
            ctx["input_path"],
        )
    return _pass("C04", check_id, ctx["input_path"], summary)


def check_c05_fuel(ctx):
    check_id = "check_fuel_source_and_vehicle"
    data = ctx["data"]
    if data is None:
        return _untestable("C05", check_id, "車両入力が無いため燃料費の原本対応を判定できません")
    fuels = _actual_records(data, "fuel")
    if not fuels:
        return _untestable("C05", check_id, "燃料の明細が無いため月別原本と車番対応を判定できません")
    bad = []
    for record in fuels:
        has_source = bool(str(record.get("source_ref") or "").strip())
        has_vehicle = bool(record.get("vehicle_id")) or any(
            item.get("vehicle_id") or item.get("vehicle_hint") for item in record.get("allocations") or []
        ) or bool(record.get("vehicle_hint"))
        if not has_source or not has_vehicle:
            bad.append(str(record.get("id") or ""))
    summary = f"fuel_records={len(fuels)} missing_source_or_vehicle={len(bad)}"
    if bad:
        return _fail(
            "C05", check_id, "SOURCE_MISMATCH",
            "燃料明細に原本参照または車番がありません: " + ",".join(bad[:12]),
            ctx["input_path"], summary,
        )
    return _pass("C05", check_id, ctx["input_path"], summary)


def check_c06_lease(ctx):
    check_id = "check_lease_amount_and_tax_note"
    data = ctx["data"]
    if data is None:
        return _untestable("C06", check_id, "車両入力が無いためリース税込注記を判定できません")
    leases = _actual_records(data, "lease")
    if not leases:
        return _untestable("C06", check_id, "リースの明細が無いため税込注記を判定できません")
    bad = []
    for record in leases:
        amount = record.get("amount")
        basis = record.get("tax_basis")
        note = str(record.get("evidence") or record.get("inclusive_note") or "")
        basis_note = str((data or {}).get("basis_note") or "")
        has_amount = amount is not None and amount != ""
        tax_ok = basis == "inclusive" or "税込" in note or "税込" in basis_note
        if not has_amount or not tax_ok:
            bad.append(str(record.get("id") or ""))
    summary = f"lease_records={len(leases)} missing_amount_or_tax_note={len(bad)}"
    if bad:
        return _fail(
            "C06", check_id, "CHECK_MISSING",
            "リース額または税込注記がありません: " + ",".join(bad[:12]),
            ctx["input_path"], summary,
        )
    return _pass("C06", check_id, ctx["input_path"], summary)


def check_c07_pnl(ctx):
    check_id = "check_vehicle_pnl_and_company_total"
    calc = ctx["calc"]
    path = ctx["calc_path"]
    data = ctx["data"]
    if not path:
        return _untestable("C07", check_id, "計算工程が無いため車両別損益の生成を判定できません")
    if calc is None:
        return _fail("C07", check_id, "CHECK_MISSING", "現行計画の計算成果物がありません", path)
    rows = list(calc.get("expected") or calc.get("rows") or [])
    if not rows and data is not None:
        try:
            from app.vehicle_profit import calculate
            rows = list(calculate(data).get("rows") or [])
        except (ValueError, TypeError, KeyError):
            rows = []
    vehicles = _vehicle_ids(data) if data is not None else sorted({r.get("vehicle") for r in rows if r.get("vehicle")})
    vehicle_rows = [r for r in rows if r.get("vehicle")]
    company_total = calc.get("company_total")
    if company_total is None and vehicle_rows:
        profits = [r.get("profit") for r in vehicle_rows]
        if all(p is not None for p in profits):
            company_total = sum(profits)
    summary = f"status={calc.get('status')} vehicle_rows={len(vehicle_rows)} company_total={company_total}"
    if calc.get("status") != "completed":
        return _fail("C07", check_id, "PROVISIONAL_ARTIFACT", "車両別月次損益が完了していません", path, summary)
    if not vehicles or len({r.get("vehicle") for r in vehicle_rows}) < len(vehicles):
        return _fail("C07", check_id, "CHECK_MISSING", "車両別月次損益が不足しています", path, summary)
    if company_total is None:
        return _fail("C07", check_id, "CHECK_MISSING", "会社合計がありません", path, summary)
    return _pass("C07", check_id, path, summary)


def check_c08_sheets(ctx):
    check_id = "check_per_vehicle_sheets"
    data = ctx["data"]
    workbook = ctx["workbook"]
    path = ctx["workbook_path"]
    if not path:
        return _untestable("C08", check_id, "Excel成果物経路が無いため車両シートを判定できません")
    if workbook is None or not getattr(workbook, "exists", lambda: False)():
        return _fail("C08", check_id, "CHECK_MISSING", "車両ごとのシートを持つExcelがありません", path)
    if data is None:
        return _untestable("C08", check_id, "車両入力が無いため対象車両数とシート数を照合できません", path)
    vehicles = _vehicle_ids(data)
    if not vehicles:
        return _fail("C08", check_id, "CHECK_MISSING", "対象車両が空です", ctx["input_path"])
    try:
        from openpyxl import load_workbook
        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            names = list(book.sheetnames)
        finally:
            book.close()
    except Exception as exc:
        return _untestable("C08", check_id, "Excelを開けないため車両シートを判定できません: " + str(exc), path)
    vehicle_sheets = [name for name in names if str(name).startswith("車両_")]
    summary = f"vehicles={len(vehicles)} vehicle_sheets={len(vehicle_sheets)} sheets={names}"
    if len(vehicle_sheets) != len(vehicles):
        return _fail(
            "C08", check_id, "CHECK_MISSING",
            f"車両シート数が対象車両数と一致しません ({len(vehicle_sheets)} != {len(vehicles)})",
            path, summary,
        )
    return _pass("C08", check_id, path, summary)


def check_c09_reconciliation(ctx):
    check_id = "check_independent_reconciliation"
    data = ctx["data"]
    summary = ctx["recon_summary"]
    rows = ctx["recon_rows"]
    if data is None:
        return _untestable("C09", check_id, "車両入力が無いため独立照合を判定できません")
    if summary is None:
        return _untestable("C09", check_id, "独立照合結果を取得できないため判定できません")
    evidence = str(summary.get("artifact") or ctx["calc_path"] or ctx["input_path"])
    text = (
        f"passed={summary.get('passed')} matched={summary.get('matched')} "
        f"mismatched={summary.get('mismatched')} unavailable={summary.get('unavailable')} rows={len(rows)}"
    )
    if not rows or summary.get("passed") is not True or any(r.get("status") != "matched" for r in rows):
        return _fail("C09", check_id, "RECONCILIATION_FAILED", "独立照合が全件matchedではありません", evidence, text)
    return _pass("C09", check_id, evidence, text)


def check_c10_unresolved(ctx):
    check_id = "check_unresolved_zero"
    data = ctx["data"]
    unresolved = ctx["unresolved"]
    if data is None:
        return _untestable("C10", check_id, "車両入力が無いため未配賦・読取失敗を判定できません")
    if unresolved is None:
        return _untestable("C10", check_id, "未解決件数を分類できないため判定できません")
    auto_issues = [i for i in ((data.get("auto_extraction") or {}).get("issues") or []) if i.get("status") != "resolved"]
    recon = ctx["recon_summary"] or {}
    unmatched = 0
    if recon:
        unmatched = int(recon.get("mismatched") or 0) + int(recon.get("unavailable") or 0) + int(recon.get("incomplete") or 0)
    total = int(unresolved.get("total") or 0)
    summary = f"unresolved={total} auto_open={len(auto_issues)} recon_unmatched={unmatched} by_class={unresolved.get('by_class')}"
    if total or unmatched:
        code = "ALLOCATION_UNRESOLVED" if (unresolved.get("by_class") or {}).get("fact_pending") else "PROVISIONAL_ARTIFACT"
        if (unresolved.get("by_class") or {}).get("missing_source"):
            code = "SOURCE_READ_FAILED"
        return _fail("C10", check_id, code, "未配賦・読取失敗・照合不能・根拠不明が残っています", ctx["input_path"], summary)
    return _pass("C10", check_id, ctx["input_path"], summary)


def check_c11_recalculation(ctx):
    check_id = "check_excel_recalculation"
    calc = ctx["calc"]
    path = ctx["calc_path"]
    if calc is None:
        return _untestable("C11", check_id, "計算成果物が無いためExcel再計算を判定できません")
    checks = calc.get("checks") or {}
    if "recalculation" not in checks:
        return _untestable("C11", check_id, "recalculation検査結果が無いため判定できません", path)
    errors = checks.get("formula_errors")
    summary = f"recalculation={checks.get('recalculation')} formula_errors={errors}"
    if checks.get("recalculation") is not True:
        return _fail("C11", check_id, "CHECK_MISSING", "Excel再計算が可能ではありません", path, summary)
    if errors not in (0, None):
        return _fail("C11", check_id, "CHECK_MISSING", "数式エラーが残っています", path, summary)
    return _pass("C11", check_id, path, summary)


def check_c12_trace(ctx):
    check_id = "check_amount_trace_to_source"
    data = ctx["data"]
    if data is None:
        return _untestable("C12", check_id, "車両入力が無いため原本追跡を判定できません")
    records = [r for r in _actual_records(data) if r.get("category") in MAJOR_TRACE_CATEGORIES]
    if not records:
        return _untestable("C12", check_id, "主要金額の明細が無いため原本追跡を判定できません")
    bad = []
    for record in records:
        if not str(record.get("source_ref") or "").strip() or not str(record.get("source_locator") or "").strip():
            bad.append(str(record.get("id") or ""))
    summary = f"major_records={len(records)} missing_trace={len(bad)}"
    if bad:
        return _fail(
            "C12", check_id, "CHECK_MISSING",
            "主要金額に原本参照または位置がありません: " + ",".join(bad[:12]),
            ctx["input_path"], summary,
        )
    return _pass("C12", check_id, ctx["input_path"], summary)


CHECKS = {
    "C01": check_c01_scope,
    "C02": check_c02_coverage,
    "C03": check_c03_categories,
    "C04": check_c04_allocation_rules,
    "C05": check_c05_fuel,
    "C06": check_c06_lease,
    "C07": check_c07_pnl,
    "C08": check_c08_sheets,
    "C09": check_c09_reconciliation,
    "C10": check_c10_unresolved,
    "C11": check_c11_recalculation,
    "C12": check_c12_trace,
}


def evaluate_vehicle_criteria(manager, project_id, criteria, contract=None):
    """Run positive checks. Structured failures override PASS for their criterion only."""
    ctx = collect_context(manager, project_id, contract)
    failures = []
    try:
        failures = list(goal_failure_records(manager, project_id) or [])
    except Exception:
        failures = []
    by_cid: dict[str, list[dict]] = {}
    for record in failures:
        for cid in record.get("criterion_ids") or []:
            by_cid.setdefault(cid, []).append(record)
    rows = []
    for item in criteria:
        cid = item.get("criterion_id")
        checker = CHECKS.get(cid)
        if checker is None:
            rows.append(_untestable(cid, "", f"{cid}の検査が未実装です", code="CHECK_UNIMPLEMENTED"))
            continue
        result = checker(ctx)
        hits = by_cid.get(cid) or []
        if hits:
            hit = hits[0]
            rows.append(_fail(
                cid, result.get("check_id") or "",
                hit.get("code") or "PROVISIONAL_ARTIFACT",
                hit.get("message") or "",
                hit.get("evidence_path") or result.get("evidence_path") or "",
                result.get("evidence_summary") or "",
            ))
            continue
        rows.append(result)
    return rows
