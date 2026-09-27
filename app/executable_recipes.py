"""Human-reviewed executable recipes: versioned procedures, fresh extraction, no amount reuse.

P2-2: RecipeSpec は AdapterSpec と同じ管理方針（provenance だけを分ける）。
human_result は由来であり適用条件ではない。金額は保存・再利用しない。
動的 exec は禁止。adapters はリポジトリ内の関数IDのみ。
"""
from __future__ import annotations

import re
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path

from app.goal_review import ReviewStore

RECIPE_STATUSES = ('verified', 'disabled', 'revoked', 'needs_revalidation')
ALLOWED_RECIPE_ADAPTERS = frozenset({
    'identity_passthrough',
    'vehicle_extract_pipeline',
    'trial_amount_map',
})


def default_applicability(**overrides):
    row = {
        'source_format_signatures': [],
        'vendor': [],
        'required_columns': [],
        'required_source_kinds': [],
        'months_pattern': '',
        'missing_policy': 'unknown',
        'no_amount_reuse': True,
        'forbidden_conditions': [],
    }
    row.update(overrides)
    row['no_amount_reuse'] = True
    return row


def default_validation_contract(**overrides):
    row = {
        'requires_source_reconciliation': True,
        'checks': ['independent_source_totals', 'no_unreadable_as_zero'],
        'fixtures': [],
    }
    row.update(overrides)
    return row


def _strip_amounts(value):
    """レシピに金額を残さない。human_result 由来の数値コピーを防ぐ。"""
    if isinstance(value, dict):
        return {
            k: _strip_amounts(v)
            for k, v in value.items()
            if k not in {'amount', 'amounts', 'copied_amounts', 'human_amounts'}
        }
    if isinstance(value, list):
        return [_strip_amounts(x) for x in value]
    return value


def revision_at_least(current, minimum) -> bool:
    if not minimum:
        return True
    if not current:
        return False

    def parts(value):
        out = []
        for token in str(value).replace('-', '.').split('.'):
            out.append(int(token) if token.isdigit() else token)
        return tuple(out)

    try:
        return parts(current) >= parts(minimum)
    except TypeError:
        return str(current) >= str(minimum)


@dataclass
class RecipeSpec:
    """設計書 第8章 RecipeSpec。金額はフィールドに持たない。"""
    recipe_id: str
    version: int = 1
    status: str = 'needs_revalidation'
    executor: str = 'vehicle-auto-v1'
    extractor_revision_min: str = ''
    adapters: list = field(default_factory=list)
    mapping_rules: list = field(default_factory=list)
    allocation_rules: list = field(default_factory=list)
    applicability: dict = field(default_factory=default_applicability)
    validation_contract: dict = field(default_factory=default_validation_contract)
    provenance: dict = field(default_factory=dict)
    expires: float = 0
    experience_id: str = ''

    def to_dict(self) -> dict:
        row = _strip_amounts(asdict(self))
        row['applicability'] = default_applicability(**(row.get('applicability') or {}))
        row['validation_contract'] = default_validation_contract(**(row.get('validation_contract') or {}))
        adapters = list(row.get('adapters') or [])
        row['adapters'] = adapters
        if not adapters and row.get('status') == 'verified':
            row['status'] = 'needs_revalidation'
        if row.get('status') not in RECIPE_STATUSES:
            row['status'] = 'needs_revalidation'
        return row


@dataclass
class MatchResult:
    matched: bool
    reasons: list = field(default_factory=list)
    recipe_id: str = ''
    version: int | str = 0

    def as_dict(self):
        return {
            'matched': self.matched,
            'reasons': list(self.reasons),
            'recipe_id': self.recipe_id,
            'version': self.version,
        }


@dataclass
class ExtractionResult:
    applied: bool
    rejected: bool
    reasons: list = field(default_factory=list)
    data: dict | None = None
    recipe_id: str = ''
    version: int | str = 0
    adapter_results: list = field(default_factory=list)
    mapping_results: list = field(default_factory=list)
    allocation_results: list = field(default_factory=list)
    source_hash: str = ''
    reuse_demonstration: bool = False
    reconciliation: dict | None = None
    rag_success: bool = False

    def envelope_fields(self):
        return {
            'recipe_applied': self.applied,
            'recipe_rejected': self.rejected,
            'recipe_id': self.recipe_id if self.applied else '',
            'recipe_version': self.version if self.applied else None,
            'recipe_rejection': list(self.reasons) if (self.rejected or not self.applied) else [],
            'recipe_rule_results': {
                'adapters': list(self.adapter_results),
                'mapping_rules': list(self.mapping_results),
                'allocation_rules': list(self.allocation_results),
            },
            'recipe_reuse_demonstration': bool(self.applied and self.reuse_demonstration),
            'recipe_learning_success': False if not self.applied else bool(self.rag_success),
            'recipe_reference': None,
        }


def normalize_recipe(recipe) -> dict:
    if isinstance(recipe, RecipeSpec):
        return recipe.to_dict()
    row = _strip_amounts(dict(recipe or {}))
    provenance = dict(row.get('provenance') or {})
    if row.get('human_result') and not provenance.get('human_result'):
        provenance['human_result'] = row['human_result']
    provenance.setdefault('human_result', '')
    provenance.setdefault('source_hashes_used_for_validation', [])
    provenance.setdefault('months_used_for_validation', [])
    row['provenance'] = provenance
    row['applicability'] = default_applicability(**(row.get('applicability') or {}))
    row['validation_contract'] = default_validation_contract(**(row.get('validation_contract') or {}))
    row.setdefault('recipe_id', row.get('experience_id') or provenance.get('human_result') or '')
    row.setdefault('version', 1)
    row.setdefault('adapters', [])
    row.setdefault('mapping_rules', [])
    row.setdefault('allocation_rules', [])
    row.setdefault('executor', 'vehicle-auto-v1')
    row.setdefault('extractor_revision_min', '')
    row.setdefault('status', 'needs_revalidation')
    row.setdefault('expires', 0)
    if not row.get('adapters') and row.get('status') == 'verified':
        row['status'] = 'needs_revalidation'
    return row


def _reason(code, message):
    return {'code': code, 'message': message}


def approval_not_revoked(memory_path, pid, recipe) -> bool:
    """当時の承認がまだ取り消されていないか。原本同一性（evidence_valid）は見ない。"""
    recipe = normalize_recipe(recipe)
    if recipe.get('status') == 'revoked':
        return False
    human = (recipe.get('provenance') or {}).get('human_result') or recipe.get('human_result')
    if not human:
        return True
    row = ReviewStore(memory_path).get(pid, 'result', human)
    if row and row.get('status') == 'revoked':
        return False
    return True


def format_signature_of(source, content=''):
    name = source.get('filename') or ''
    suffix = Path(name).suffix.lower()
    if suffix in {'.xlsx', '.xlsm'}:
        return 'excel:workbook'
    if suffix == '.csv':
        return 'csv:table'
    if suffix == '.pdf':
        from app.vehicle_auto import source_vendor
        vendor = source_vendor(source, content)
        return 'pdf:' + (vendor or 'document')
    return 'file:' + (suffix.lstrip('.') or 'unknown')


def source_kind_of(source, content=''):
    name = source.get('filename') or ''
    suffix = Path(name).suffix.lower()
    if '給与' in name or '支給' in name:
        return 'payroll'
    if suffix in {'.xlsx', '.xlsm'}:
        return 'master'
    if suffix == '.csv' or suffix == '.pdf' or any(token in name for token in ('燃料', '請求', '宇佐美', 'ENEOS', 'Wing')):
        return 'fuel'
    return 'unknown'


def columns_of(source, raw=None, content=''):
    suffix = Path(source.get('filename') or '').suffix.lower()
    if suffix == '.csv':
        header = next((line for line in (content or '').splitlines() if line.strip()), '')
        return [c.strip() for c in header.split(',') if c.strip()]
    if suffix in {'.xlsx', '.xlsm'} and raw:
        try:
            import io
            from openpyxl import load_workbook
            from app.vehicle_auto import compact
            book = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            labels = []
            sheet = next(iter(book), None)
            if sheet is not None:
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i >= 10:
                        break
                    for cell in row:
                        packed = compact(cell)
                        if packed:
                            labels.append(packed)
            book.close()
            return list(dict.fromkeys(labels))
        except Exception:
            return []
    return []


def build_source_snapshot(docs, sources, months, extractor_revision='', missing_policy='unknown'):
    """match_recipe 用の原本スナップショット。金額は載せない。"""
    vendors, signatures, kinds, columns, hashes, filenames = [], [], [], [], [], []
    from app.vehicle_auto import source_vendor
    for item in docs or []:
        source = item[0] if isinstance(item, (list, tuple)) else item
        raw = item[1] if isinstance(item, (list, tuple)) and len(item) > 1 else None
        content = item[2] if isinstance(item, (list, tuple)) and len(item) > 2 else ''
        filenames.append(source.get('filename') or '')
        vendor = source_vendor(source, content or '')
        if vendor:
            vendors.append(vendor)
        signatures.append(format_signature_of(source, content or ''))
        kind = source_kind_of(source, content or '')
        if kind != 'unknown':
            kinds.append(kind)
        columns.extend(columns_of(source, raw, content or ''))
        if source.get('sha256'):
            hashes.append(source['sha256'])
    for source in sources or []:
        if source.get('sha256'):
            hashes.append(source['sha256'])
        if source.get('filename') and source['filename'] not in filenames:
            filenames.append(source['filename'])
            signatures.append(format_signature_of(source))
            kind = source_kind_of(source)
            if kind != 'unknown':
                kinds.append(kind)
            vendor = source_vendor(source)
            if vendor:
                vendors.append(vendor)
    return {
        'extractor_revision': extractor_revision,
        'vendors': sorted(set(vendors)),
        'format_signatures': sorted(set(signatures)),
        'columns': list(dict.fromkeys(columns)),
        'source_kinds': sorted(set(kinds)),
        'months': list(months or []),
        'source_hashes': list(dict.fromkeys(hashes)),
        'missing_policy': missing_policy,
        'filenames': filenames,
        'sources': list(sources or []),
    }


def is_reuse_demonstration(recipe, source_hash, months) -> bool:
    """同一 hash の再処理は再利用実証に数えない。別月かつ別 source_hash が必要。"""
    recipe = normalize_recipe(recipe)
    provenance = recipe.get('provenance') or {}
    used_hashes = set(provenance.get('source_hashes_used_for_validation') or [])
    used_months = set(provenance.get('months_used_for_validation') or [])
    if not source_hash or source_hash in used_hashes:
        return False
    new_months = set(months or [])
    if not new_months or new_months <= used_months:
        return False
    return True


def recipe_learning_success(applied, recon_summary, human_approved=False) -> bool:
    """レシピ参照だけでは RAG 成功にしない。独立照合と人間確認が必要。"""
    if not applied:
        return False
    if not recon_summary or recon_summary.get('passed') is not True:
        return False
    if not human_approved:
        return False
    return True


def match_recipe(recipe, source_snapshot) -> MatchResult:
    recipe = normalize_recipe(recipe)
    snapshot = dict(source_snapshot or {})
    recipe_id = recipe.get('recipe_id') or ''
    version = recipe.get('version') or 1
    reasons = []
    status = recipe.get('status')
    if status == 'revoked':
        reasons.append(_reason('recipe_revoked', '取り消された手順は選択しません'))
    elif status == 'disabled':
        reasons.append(_reason('recipe_disabled', '無効化された手順は選択しません'))
    elif status != 'verified':
        reasons.append(_reason('not_verified', 'verified ではない手順は適用しません'))
    if not recipe.get('adapters'):
        reasons.append(_reason('no_adapters', 'adapters の無い手順は verified として適用しません'))
    if recipe.get('executor') != 'vehicle-auto-v1':
        reasons.append(_reason('executor_mismatch', 'executor が vehicle-auto-v1 ではありません'))
    if not revision_at_least(snapshot.get('extractor_revision') or '', recipe.get('extractor_revision_min') or ''):
        reasons.append(_reason('extractor_revision_too_old', 'extractor_revision_min を満たしません'))
    appl = recipe.get('applicability') or {}
    vendors_needed = [v for v in (appl.get('vendor') or []) if v]
    vendors_have = set(snapshot.get('vendors') or [])
    if vendors_needed and not set(vendors_needed) <= vendors_have:
        reasons.append(_reason('vendor_mismatch', 'vendor が手順の適用条件と一致しません'))
    signatures_needed = [s for s in (appl.get('source_format_signatures') or []) if s]
    signatures_have = set(snapshot.get('format_signatures') or [])
    if signatures_needed and not set(signatures_needed) <= signatures_have:
        reasons.append(_reason('format_signature_mismatch', 'source format signature が一致しません'))
    required_columns = [c for c in (appl.get('required_columns') or []) if c]
    have_columns = set(snapshot.get('columns') or [])
    missing_columns = [c for c in required_columns if c not in have_columns]
    if missing_columns:
        reasons.append(_reason('required_columns_missing', '必須列が欠落しています: ' + ','.join(missing_columns)))
    required_kinds = [k for k in (appl.get('required_source_kinds') or []) if k]
    have_kinds = set(snapshot.get('source_kinds') or [])
    missing_kinds = [k for k in required_kinds if k not in have_kinds]
    if missing_kinds:
        reasons.append(_reason(
            'required_source_kinds_missing',
            '必須原本種別がありません: ' + ','.join(missing_kinds),
        ))
    pattern = appl.get('months_pattern') or ''
    if pattern:
        for mon in snapshot.get('months') or []:
            if not re.fullmatch(pattern, str(mon) or ''):
                reasons.append(_reason('months_pattern_mismatch', '対象月が months_pattern に一致しません'))
                break
    policy = appl.get('missing_policy')
    if policy in {'zero', 'unknown'} and snapshot.get('missing_policy') not in (None, '', policy):
        reasons.append(_reason('missing_policy_mismatch', '不足0円方針が手順と一致しません'))
    blob = ' '.join(snapshot.get('filenames') or []) + ' ' + ' '.join(snapshot.get('columns') or [])
    for cond in appl.get('forbidden_conditions') or []:
        if cond and str(cond) in blob:
            reasons.append(_reason('forbidden_condition', '禁止条件に該当します: ' + str(cond)))
            break
    if appl.get('no_amount_reuse') is False:
        reasons.append(_reason('amount_reuse_forbidden', '金額再利用は禁止です'))
    return MatchResult(matched=not reasons, reasons=reasons, recipe_id=recipe_id, version=version)


def _adapter_function_id(adapter) -> str:
    return str(adapter.get('function_id') or adapter.get('adapter_id') or '')


def execute_recipe(recipe, docs, months, *, zero=False, source_hash='', snapshot=None) -> ExtractionResult:
    """新原本で必ず extract する。recipe から金額をコピーしない。"""
    from app.vehicle_auto import apply_adopted_function, extract
    from app.vehicle_workflow import source_reconciliation_report

    recipe = normalize_recipe(recipe)
    recipe_id = recipe.get('recipe_id') or ''
    version = recipe.get('version') or 1
    match = match_recipe(recipe, snapshot or {})
    if not match.matched:
        return ExtractionResult(
            applied=False, rejected=True, reasons=list(match.reasons),
            recipe_id=recipe_id, version=version, source_hash=source_hash,
        )
    if recipe.get('mapping_rules'):
        return ExtractionResult(
            applied=False, rejected=True,
            reasons=[_reason('mapping_rules_unsupported', 'mapping_rules が空でないため適用を拒否します')],
            recipe_id=recipe_id, version=version, source_hash=source_hash,
        )
    adapter_results = []
    for adapter in recipe.get('adapters') or []:
        function_id = _adapter_function_id(adapter)
        params = dict(adapter.get('params') or {})
        if function_id not in ALLOWED_RECIPE_ADAPTERS:
            return ExtractionResult(
                applied=False, rejected=True,
                reasons=[_reason('unknown_adapter', '未登録のアダプター関数IDです（動的execは禁止）: ' + function_id)],
                recipe_id=recipe_id, version=version, source_hash=source_hash,
                adapter_results=adapter_results,
            )
        if 'amount' in params or 'amounts' in params:
            return ExtractionResult(
                applied=False, rejected=True,
                reasons=[_reason('amount_reuse_forbidden', 'adapter params からの金額再利用は禁止です')],
                recipe_id=recipe_id, version=version, source_hash=source_hash,
                adapter_results=adapter_results,
            )
        adapter_results.append({
            'adapter_id': adapter.get('adapter_id') or function_id,
            'version': adapter.get('version') or '',
            'function_id': function_id,
            'status': 'selected',
        })
    data = extract(docs, months, zero)
    for adapter in recipe.get('adapters') or []:
        function_id = _adapter_function_id(adapter)
        params = {
            k: v for k, v in dict(adapter.get('params') or {}).items()
            if k not in {'amount', 'amounts', 'copied_amounts', 'human_amounts'}
        }
        data = apply_adopted_function(function_id, data, params)
        for row in adapter_results:
            if row.get('function_id') == function_id:
                row['status'] = 'applied'
    mapping_results = [{'status': 'reserved', 'rules': recipe.get('mapping_rules') or []}]
    allocation_results = [{'status': 'reserved', 'rules': recipe.get('allocation_rules') or []}]
    _, recon_summary, _ = source_reconciliation_report(data, (snapshot or {}).get('sources'))
    contract = recipe.get('validation_contract') or {}
    rag = False
    if contract.get('requires_source_reconciliation') and recon_summary.get('passed') is not True:
        rag = False
    reuse = is_reuse_demonstration(recipe, source_hash, months)
    return ExtractionResult(
        applied=True, rejected=False, reasons=[], data=data,
        recipe_id=recipe_id, version=version, adapter_results=adapter_results,
        mapping_results=mapping_results, allocation_results=allocation_results,
        source_hash=source_hash, reuse_demonstration=reuse,
        reconciliation=recon_summary, rag_success=rag,
    )


def _selectable(recipe) -> bool:
    recipe = normalize_recipe(recipe)
    if recipe.get('status') != 'verified':
        return False
    if recipe.get('executor') != 'vehicle-auto-v1':
        return False
    if recipe.get('expires') and recipe['expires'] <= time.time():
        return False
    if not recipe.get('adapters'):
        return False
    return True


def register(manager, pid, signature, result):
    from app.vehicle_workflow import applicable, load_input, sources as list_sources
    from app.vehicle_auto import REVISION, source_vendor

    mission = manager.memory.get_mission(pid)
    envelope = load_input(manager, pid) if applicable(mission) else None
    if not envelope or envelope.get('mode') != 'vehicle-auto-v1':
        return None
    data = envelope.get('data') or {}
    snapshot = list_sources(manager, pid)
    adapters = list(envelope.get('recipe_adapters') or [])
    if not adapters:
        adapters = [{'adapter_id': 'vehicle_extract_pipeline', 'version': '1', 'params': {}, 'function_id': 'vehicle_extract_pipeline'}]
    vendors = sorted({source_vendor(s) for s in snapshot if source_vendor(s)})
    signatures = sorted({format_signature_of(s) for s in snapshot})
    kinds = sorted({k for s in snapshot if (k := source_kind_of(s)) != 'unknown'})
    hashes = [s.get('sha256') for s in snapshot if s.get('sha256')]
    spec = RecipeSpec(
        recipe_id='recipe-' + str(signature)[:16],
        version=1,
        status='verified' if adapters else 'needs_revalidation',
        executor='vehicle-auto-v1',
        extractor_revision_min=str(envelope.get('extractor_revision') or REVISION),
        adapters=adapters,
        mapping_rules=[],
        allocation_rules=list(data.get('allocation_rules') or envelope.get('assignments') or []),
        applicability=default_applicability(
            source_format_signatures=signatures,
            vendor=vendors,
            required_columns=[],
            required_source_kinds=kinds,
            months_pattern='',
            missing_policy=data.get('missing_policy') or 'unknown',
        ),
        validation_contract=default_validation_contract(),
        provenance={
            'human_result': signature,
            'source_hashes_used_for_validation': hashes,
            'months_used_for_validation': list(data.get('months') or []),
        },
        expires=result.get('expires') or 0,
        experience_id=result.get('experience_id') or '',
    )
    payload = spec.to_dict()
    ReviewStore(manager.memory.path).put(pid, 'executable_recipe', signature, payload)
    return payload


def select(manager, pid, source_snapshot=None):
    """applicability と extractor_revision_min で判定する。evidence_valid は再利用条件に使わない。"""
    for row in verified_recipes(manager, pid):
        if source_snapshot is not None:
            match = match_recipe(row, source_snapshot)
            if not match.matched:
                continue
        return row
    return None


def verified_recipes(manager, pid):
    store = ReviewStore(manager.memory.path)
    rows = []
    for _, payload in store.list(pid, 'executable_recipe'):
        row = normalize_recipe(payload)
        if not _selectable(row):
            continue
        if not approval_not_revoked(manager.memory.path, pid, row):
            continue
        rows.append(row)
    return list(reversed(rows))


def apply_during_prepare(manager, pid, docs, months, zero, snapshot, extractor_revision, missing_policy, ocr_adoptions=None):
    """prepare から呼ぶ。不一致は明示拒否したあと、未適用の通常抽出へ進む。silent fallback 禁止。"""
    from app.vehicle_auto import extract
    from app.vehicle_workflow import digest, source_reconciliation_report

    view = build_source_snapshot(docs, snapshot, months, extractor_revision, missing_policy)
    recipes = verified_recipes(manager, pid)
    if not recipes:
        data = extract(docs, months, zero, ocr_adoptions=ocr_adoptions)
        return data, {
            'recipe_applied': False,
            'recipe_rejected': False,
            'recipe_id': '',
            'recipe_version': None,
            'recipe_rejection': [],
            'recipe_rule_results': {},
            'recipe_reuse_demonstration': False,
            'recipe_learning_success': False,
            'recipe_reference': None,
        }
    last_mismatch = None
    for recipe in recipes:
        match = match_recipe(recipe, view)
        if match.matched:
            result = execute_recipe(
                recipe, docs, months, zero=zero,
                source_hash=digest(snapshot), snapshot=view,
            )
            if result.rejected or not result.applied or result.data is None:
                data = extract(docs, months, zero, ocr_adoptions=ocr_adoptions)
                fields = result.envelope_fields()
                fields['recipe_applied'] = False
                fields['recipe_rejected'] = True
                fields['recipe_reuse_demonstration'] = False
                fields['recipe_learning_success'] = False
                return data, fields
            _, recon, _ = source_reconciliation_report(result.data, snapshot)
            fields = result.envelope_fields()
            fields['recipe_learning_success'] = recipe_learning_success(True, recon, False)
            fields['recipe_reconciliation'] = recon
            return result.data, fields
        last_mismatch = (recipe, match)
    recipe, match = last_mismatch
    data = extract(docs, months, zero, ocr_adoptions=ocr_adoptions)
    return data, {
        'recipe_applied': False,
        'recipe_rejected': True,
        'recipe_id': recipe.get('recipe_id') or '',
        'recipe_version': recipe.get('version'),
        'recipe_rejection': list(match.reasons),
        'recipe_rule_results': {'match': match.as_dict()},
        'recipe_reuse_demonstration': False,
        'recipe_learning_success': False,
        'recipe_reference': None,
    }
