"""Typed local vehicle accounting workflow. Documents cannot satisfy calculation goals."""
from __future__ import annotations
import asyncio
import copy
import hashlib
import io
import json
import re
import unicodedata
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from app.structured_planning import SCHEMA, contract_of
from app.vehicle_profit import build_workbook, calculate, number

CATEGORIES = ['revenue', 'payroll', 'insurance', 'fuel', 'toll', 'lease', 'other']
INPUT_PATH = 'vehicle_profit/input.json'
APPLY_THIS_MONTH = 'this_month'
APPLY_DATE_RANGE = 'date_range'
INSURANCE_EMPLOYER_COLUMNS = ('社保会社負担', '社会保険会社負担', '事業主負担額', '会社負担社会保険')
INSURANCE_PERSONAL_REJECT = (
    '社会保険料', '健康保険料', '厚生年金', '厚生年金保険料', '介護保険料', '雇用保険料',
    '社保個人負担', '個人負担社会保険', '本人負担社会保険', '健康保険料個人負担',
    '社会保険料個人負担', '個人負担', '本人負担', '健康保険料の個人負担',
    '控除社会保険', '社会保険控除', '社会保険料控除', '健康保険料控除',
)


class AllocationRuleError(ValueError):
    """配賦ルール検証失敗。API は 422 で返す。"""
    http_status = 422


def _month_or_none(value):
    if value in (None, ''):
        return None
    text = str(value)
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', text):
        raise AllocationRuleError('有効期間は YYYY-MM で指定してください')
    return text


def _ratio_total(allocations):
    total = Decimal('0')
    for item in allocations:
        ratio = Decimal(str(item.get('ratio')))
        if ratio <= 0:
            raise AllocationRuleError('配賦比率は0より大きい値が必要です')
        total += ratio
    return total


def allocation_periods_overlap(left_from, left_to, right_from, right_to):
    a0 = left_from or '0000-01'
    a1 = left_to or '9999-12'
    b0 = right_from or '0000-01'
    b1 = right_to or '9999-12'
    return a0 <= b1 and b0 <= a1


def subject_key_of(rule):
    key = rule.get('subject_key') or {}
    return (rule.get('subject_type') or 'employee', key.get('company') or '', key.get('employee') or key.get('subject') or '')


def make_allocation_rule(*, subject_type='employee', subject_key=None, effective_from=None, effective_to=None,
                         allocations=None, basis_type='manual_business_fact', basis_evidence='', reason='',
                         source_hash='', created_by='human', decision_id=None, status='active'):
    allocations = list(allocations or [])
    rule = {
        'id': 'allocation-rule-' + digest([subject_type, subject_key, effective_from, effective_to, allocations, reason])[:16],
        'subject_type': subject_type,
        'subject_key': dict(subject_key or {}),
        'effective_from': effective_from,
        'effective_to': effective_to,
        'allocations': allocations,
        'basis_type': basis_type,
        'basis_evidence': basis_evidence,
        'reason': reason,
        'source_hash': source_hash,
        'created_by': created_by,
        'decision_id': decision_id,
        'status': status,
    }
    return rule


def validate_allocation_rule(rule, vehicles, existing_rules=None, records=None):
    """ratio合計1、会社整合、期間重複を検査する。失敗は 422。"""
    if not isinstance(rule, dict):
        raise AllocationRuleError('配賦ルールが不正です')
    allocations = list(rule.get('allocations') or [])
    if not allocations:
        raise AllocationRuleError('配賦先がありません')
    try:
        total = _ratio_total(allocations)
    except AllocationRuleError:
        raise
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise AllocationRuleError('配賦比率が不正です') from exc
    if total != Decimal('1'):
        raise AllocationRuleError('配賦比率の合計は1である必要があります')
    ids = [v['id'] for v in vehicles or []]
    by_id = {v['id']: v for v in vehicles or []}
    seen = set()
    record_companies = {r.get('company') for r in (records or []) if r.get('company')}
    for item in allocations:
        vid = item.get('vehicle_id')
        if vid not in ids or vid in seen:
            raise AllocationRuleError('配賦先の車両が対象集合にないか重複しています')
        seen.add(vid)
        company = by_id[vid].get('company')
        if record_companies and company and any(c not in ('', '未特定', company) for c in record_companies):
            raise AllocationRuleError('配賦先の会社が明細と一致しません')
    effective_from = _month_or_none(rule.get('effective_from'))
    effective_to = _month_or_none(rule.get('effective_to'))
    if effective_from and effective_to and effective_from > effective_to:
        raise AllocationRuleError('有効期間の開始が終了より後です')
    subject = subject_key_of(rule)
    for other in existing_rules or []:
        if other.get('status') not in (None, 'active'):
            continue
        if other.get('id') and other.get('id') == rule.get('id'):
            continue
        if subject_key_of(other) != subject:
            continue
        if allocation_periods_overlap(effective_from, effective_to, other.get('effective_from'), other.get('effective_to')):
            raise AllocationRuleError('同一対象の配賦期間が重複しています')
    return True


def insurance_column_rejected(label):
    packed = re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(label or '')).strip())
    return packed in INSURANCE_PERSONAL_REJECT


def employer_insurance_column(labels):
    packed = [re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(x or '')).strip()) for x in labels]
    for i, name in enumerate(packed):
        if name in INSURANCE_EMPLOYER_COLUMNS:
            return i
    return None


def digest(value):
    return hashlib.sha256(value if isinstance(value, bytes) else json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def applicable(mission):
    text = mission.get('goal', '')
    return bool(re.search(r'車両|車番', text) and re.search(r'損益|収益', text) and re.search(r'Excel|エクセル|xlsx', text, re.I))


def requirements_hash(mission):
    return digest({**{k:mission.get(k) for k in ['goal', 'success_criteria', 'constraints_text']}, 'instructions':[x.get('message') for x in mission.get('instruction_messages',[]) if x.get('kind')=='mission_instruction_user']})


def requested_months(mission):
    text=unicodedata.normalize('NFKC', mission.get('success_criteria','')+'\n'+mission.get('goal',''))
    match=re.search(r'(20\d{2})年\s*(\d{1,2})月?\s*(?:から|〜|～|~|－|-)\s*(\d{1,2})月', text)
    if not match: return []
    year,start,end=map(int,match.groups())
    return [f'{year}-{m:02d}' for m in range(start,end+1)] if 1<=start<=end<=12 else []


def make_plan(mission, criteria, source_ids):
    version=mission['plan_version']+1
    base=f'result/vehicle/v{version}'
    common={'schema':SCHEMA,'external_actions':'approval_required','workflow_version':version,
            'requirements_hash':requirements_hash(mission),'executor_version':'vehicle-auto-v1','source_refs':['context:'+x for x in source_ids],
            'criterion_ids':[],'inputs':[],'months':requested_months(mission)}
    def task(key,title,kind,outputs,parents,ids):
        c=dict(common,execution_kind=kind,outputs=outputs,criterion_ids=ids,
               criterion='\n'.join(criteria) if ids else title)
        return dict(task_key=key,title=title,mode='local',depends_on=parents,
                    description=title+'。実資料と計算証拠を検査する。文書作成を実計算の完了として扱わない。',
                    acceptance_criteria=json.dumps(c,ensure_ascii=False))
    tasks=[task('SC00','原本の読取・入力資料の整理','vehicle_sources',
                [{'path':base+'/sources.json'}],[],[]),
           task('vehicle_extract','原本から明細・車両・月を自動対応','vehicle_extract',[{'path':base+'/normalization.json'}],['SC00'],[]),
           task('vehicle_calculate','車両・月別の損益計算とExcel生成','vehicle_calculate',
                [{'path':base+'/calculation.json'},{'path':base+'/profit.xlsx'}],['vehicle_extract'],
                [f'SC{i:02d}' for i in range(1,len(criteria)+1)]),
           task('final_verification','目標・原本・Excelの最終照合','vehicle_verify',
                [{'path':base+'/verification.json'}],['vehicle_calculate'],[])]
    c=contract_of(tasks[-1]);c['final_verification']=True;tasks[-1]['acceptance_criteria']=json.dumps(c,ensure_ascii=False)
    return {'summary':'車両別損益の実行計画。原本読取→入力・対応確認→月別計算→Excel再計算→独立照合。未確定項目があれば途中成果を保持して確認待ちにする。\n'+json.dumps({'criteria':criteria,'goal':mission['goal'],'months':common['months'],'version':version},ensure_ascii=False),'tasks':tasks}


def resolve(manager,pid,path,exists=False):
    project=manager.memory.get_project(pid)
    return manager.workspace.resolve_file(project.get('workspace_path',''),pid,path,must_exist=exists)[2]


def sources(manager,pid):
    from app.context_files import extraction_quality
    result=[]
    for item in manager.memory.list_context_files(pid,include_content=True):
        if item.get('source')=='memo':continue
        full=manager.memory.get_context_file(pid,item['id']) or item
        raw=full.get('original_data')
        text=str(full.get('content',''))
        result.append({'id':item['id'],'filename':item['filename'],'sha256':digest(bytes(raw)) if raw else digest(text.encode()),
                       'original_available':bool(raw),'extraction_sha256':digest(text.encode()),'quality':extraction_quality(text)})
    return sorted(result,key=lambda x:x['id'])


def write_json(path,value):
    path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix(path.suffix+'.tmp');temp.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8');temp.replace(path)


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def prepare_sources(manager,pid,task,mission):
    from app.context_files import extract_context_file
    items=sources(manager,pid);contract=contract_of(task);target=resolve(manager,pid,contract['outputs'][0]['path']);base=target.parent
    for item in items:
        full=manager.memory.get_context_file(pid,item['id'])
        raw=full.get('original_data');text=full.get('content','');suffix=Path(item['filename']).suffix.lower()
        # Extract cell addresses, not invented normalized amounts. Preserve raw source.
        if raw and suffix in {'.xlsx','.xlsm'}:
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
            cells=[];truncated=False
            for sheet in book:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            cells.append({'sheet':sheet.title,'cell':cell.coordinate,'value':str(cell.value)})
                            if len(cells)>=30000:truncated=True;break
                    if truncated:break
                if truncated:break
            book.close()
            write_json(base/'extracted'/(item['id']+'.json'),{'source':item,'cells':cells,'truncated':truncated})
        elif raw:
            extracted=extract_context_file(item['filename'],bytes(raw));text=extracted.content
        text_path=base/'extracted'/(item['id']+'.txt');text_path.parent.mkdir(parents=True,exist_ok=True)
        text_path.write_text(str(text)[:200000],encoding='utf-8')
    result={'version':mission['plan_version'],'sources':items,'requirements_hash':requirements_hash(mission)}
    write_json(target,result)
    (base/'input_guide.md').write_text('登録原本から明細と車両・月対応を自動抽出します。未解決事項はExcelの「不足と未配賦」で確認できます。全量JSONの転記は不要です。',encoding='utf-8')
    return '登録原本を保存しました。続いて月次明細を自動抽出します。'



def match_vehicle(hint, vehicles, company=None):
    hint=unicodedata.normalize('NFKC', str(hint)).strip()
    if not re.fullmatch(r'\d{3,4}',hint):raise ValueError('車番候補は下3桁または4桁で指定してください')
    matches=[]
    for vehicle in vehicles:
        if company and vehicle.get('company')!=company:continue
        tail=re.search(r'(\d+)$',unicodedata.normalize('NFKC',vehicle['id']))
        if tail and tail[1].zfill(4).endswith(hint):matches.append(vehicle['id'])
    if len(matches)!=1:raise ValueError('車番 '+hint+': 対応候補が'+str(len(matches))+'件です。会社と完全な車両IDを確認してください')
    return matches[0]


def normalize_input(data, months):
    data=copy.deepcopy(data)
    if months and data.get('months')!=months:raise ValueError('対象年月が目標と一致しません')
    for record in data.get('records',[]):
        for allocation in record.get('allocations',[]):
            if not allocation.get('vehicle_id') and allocation.get('vehicle_hint'):
                allocation['vehicle_id']=match_vehicle(allocation['vehicle_hint'],data.get('vehicles',[]),allocation.get('company'))
    data['required_categories']=CATEGORIES
    calculate(data)  # Validate existing arithmetic and IDs before filling gaps.
    return data


def _is_policy_or_gap_record(record):
    if record.get('quality') in {'missing', 'assumed_zero'}:
        return True
    return str(record.get('source_ref') or '').startswith('policy:')


def _dispositions_complete(data, snapshot):
    if snapshot is None:
        return True
    known = {'context:' + s['id'] for s in snapshot}
    dispositions = {x.get('source_ref'): x for x in data.get('source_dispositions', [])}
    for ref in known:
        d = dispositions.get(ref, {})
        if d.get('status') not in {'included', 'excluded'} or not str(d.get('evidence', '')).strip():
            return False
    return True


def source_reconciliation_report(data, snapshot=None):
    """独立原本照合の詳細結果。空配列や issues 空では合格にしない。"""
    from app.vehicle_reconciliation import reconcile_sources
    records = [r for r in data.get('records', []) if not _is_policy_or_gap_record(r)]
    controls = list(data.get('source_controls') or [])
    dispositions = list(data.get('source_dispositions') or [])
    excluded = list(data.get('excluded_records') or [])
    read_issues = [
        i for i in (data.get('auto_extraction') or {}).get('issues', [])
        if i.get('kind') == 'read' and i.get('status') != 'resolved'
    ]
    rows = reconcile_sources(records, controls, dispositions, excluded, read_issues)
    counts = {'matched': 0, 'mismatched': 0, 'unavailable': 0, 'incomplete': 0}
    issues = []
    for row in rows:
        status = row.get('status')
        if status in counts:
            counts[status] += 1
        else:
            counts['incomplete'] += 1
        key = '/'.join(str(row.get(k) or '') for k in ('source_ref', 'month', 'category', 'tax_basis'))
        if status == 'mismatched':
            issues.append(key + ': 原本合計と明細合計が未照合または不一致です')
        elif status == 'unavailable':
            issues.append(key + ': 原本統制値が取得不能のため確定できません')
        elif status != 'matched':
            issues.append(key + ': 読取失敗・未根拠除外または重複統制のため照合未完了です')
    passed = bool(rows) and all(r.get('status') == 'matched' for r in rows)
    if not _dispositions_complete(data, snapshot):
        passed = False
    if not rows:
        issues.append('原本統制による独立照合結果がありません')
        passed = False
    summary = {
        'passed': passed is True,
        'matched': counts['matched'],
        'mismatched': counts['mismatched'],
        'unavailable': counts['unavailable'],
        'incomplete': counts['incomplete'],
        'artifact': '',
    }
    return issues, summary, rows


def _goal_failure(code, criterion_ids, message, evidence_path=''):
    return {
        'code': code,
        'criterion_ids': list(criterion_ids),
        'evidence_path': evidence_path,
        'message': message,
    }


def validate_evidence_records(data, snapshot, confirmed):
    """Structured evidence failures. criterion_ids are assigned at each check site."""
    records = []
    known = {'context:' + s['id']: s for s in snapshot}
    auto = bool(data.get('auto_extraction'))
    if not auto:
        if not confirmed:
            records.append(_goal_failure('CHECK_MISSING', ['C01'], '入力内容・対象車両・原本照合の確認が未完了です'))
        if not str(data.get('scope_evidence', '')).strip():
            records.append(_goal_failure('CHECK_MISSING', ['C01'], '対象車両の網羅性と対象期間の根拠が必要です'))
        dispositions = {x.get('source_ref'): x for x in data.get('source_dispositions', [])}
        for ref in known:
            d = dispositions.get(ref, {})
            if d.get('status') not in {'included', 'excluded'} or not str(d.get('evidence', '')).strip():
                records.append(_goal_failure('PROVISIONAL_ARTIFACT', ['C10'], ref + ': 原本の採用・除外理由が未確認です'))
        seen = set()
        for r in data.get('records', []):
            if _is_policy_or_gap_record(r):
                continue
            ref = r.get('source_ref')
            source = known.get(ref)
            if not source:
                records.append(_goal_failure('CHECK_MISSING', ['C12'], str(r.get('id')) + ': 登録原本参照が不正です'))
                continue
            if not str(r.get('source_locator', '')).strip():
                records.append(_goal_failure('CHECK_MISSING', ['C12'], r['id'] + ': 原本のページ・セル位置が必要です'))
            if dispositions.get(ref, {}).get('status') != 'included':
                records.append(_goal_failure('PROVISIONAL_ARTIFACT', ['C10'], r['id'] + ': 除外原本から明細が作られています'))
            if source['quality'] != 'readable' and not str(dispositions.get(ref, {}).get('readback_evidence', '')).strip():
                records.append(_goal_failure('SOURCE_READ_FAILED', ['C10'], ref + ': 文字化け・読取不足の原本照合結果が必要です'))
            identity = (ref, r.get('source_locator'), r['month'], r['category'])
            if identity in seen:
                records.append(_goal_failure('SOURCE_MISMATCH', ['C09'], r['id'] + ': 同じ原本位置・月・費目の二重計上です'))
            seen.add(identity)
    recon_issues, _, recon_rows = source_reconciliation_report(data, snapshot)
    for issue, row in zip(recon_issues, recon_rows or []):
        status = row.get('status')
        if status == 'mismatched':
            code = 'SOURCE_MISMATCH'
        elif status == 'unavailable':
            code = 'RECONCILIATION_FAILED'
        else:
            code = 'SOURCE_READ_FAILED'
        records.append(_goal_failure(code, ['C09'], issue))
    extra_issues = recon_issues[len(recon_rows or []):]
    for issue in extra_issues:
        records.append(_goal_failure('RECONCILIATION_FAILED', ['C09'], issue))
    if auto:
        from app.vehicle_auto import evidence_issues
        for message in evidence_issues(data, snapshot):
            code = 'PROVISIONAL_ARTIFACT'
            cids = ['C10']
            if '原本参照が不正' in message:
                code, cids = 'CHECK_MISSING', ['C12']
            elif '読み取れません' in message or '読取' in message:
                code, cids = 'SOURCE_READ_FAILED', ['C10']
            records.append(_goal_failure(code, cids, message))
    unique = []
    seen_msg = set()
    for item in records:
        if item['message'] in seen_msg:
            continue
        seen_msg.add(item['message'])
        unique.append(item)
    return unique


def validate_evidence(data, snapshot, confirmed):
    return [item['message'] for item in validate_evidence_records(data, snapshot, confirmed)]


def load_input(manager,pid):
    path=resolve(manager,pid,INPUT_PATH)
    return read_json(path) if path.exists() else None


def classify_unresolved(data):
    """Classify leftover allocation/read issues. Never invent amounts."""
    issues = list((data.get('auto_extraction') or {}).get('issues') or [])
    unresolved = [i for i in issues if i.get('status') != 'resolved']
    buckets = {
        'missing_source': [],
        'mapping_missing': [],
        'fact_pending': [],
        'control_unavailable': [],
        'known_defect': [],
        'already_answered_applicable': [],
    }
    for issue in unresolved:
        kind = issue.get('kind')
        if kind == 'read':
            buckets['missing_source'].append(issue)
        elif kind == 'adapter':
            buckets['mapping_missing'].append(issue)
        elif kind == 'allocation':
            if issue.get('answered'):
                buckets['already_answered_applicable'].append(issue)
            else:
                buckets['fact_pending'].append(issue)
        elif issue.get('development'):
            buckets['known_defect'].append(issue)
        else:
            buckets['fact_pending'].append(issue)
    for record in data.get('records') or []:
        if record.get('quality') in {'missing', 'assumed_zero'}:
            continue
        if not record.get('allocations'):
            buckets['fact_pending'].append({'kind': 'allocation', 'record_id': record.get('id')})
    for control in data.get('source_controls') or []:
        if control.get('status') in {'unavailable', 'missing'}:
            buckets['control_unavailable'].append(control)
    questions = []
    for item in buckets['fact_pending']:
        questions.append({
            'subject': item.get('subject') or item.get('record_id') or item.get('id') or '',
            'kind': item.get('kind') or 'allocation',
            'hint': item.get('message') or item.get('reason') or '配賦または業務事実の確認が必要です',
        })
    return {
        'total': sum(len(v) for k, v in buckets.items() if k != 'already_answered_applicable'),
        'by_class': {key: len(value) for key, value in buckets.items()},
        'questions': questions[:50],
    }


def goal_failure_records(manager, pid):
    """Structured failures. Each check site assigns criterion_ids and code."""
    mission = manager.memory.get_mission(pid)
    if not applicable(mission):
        return []
    task = next((t for t in mission['tasks'] if (contract_of(t) or {}).get('execution_kind') == 'vehicle_calculate'), None)
    if not task:
        return [_goal_failure('CHECK_MISSING', ['C07'], '車両別・月別の実計算工程がありません。計画書だけでは目標達成にできません')]
    c = contract_of(task)
    calc_path = c['outputs'][0]['path']
    path = resolve(manager, pid, calc_path)
    if not path.exists():
        return [_goal_failure('CHECK_MISSING', ['C07'], '現行計画の計算・Excel検証記録がありません', calc_path)]
    try:
        result = read_json(path)
        envelope = load_input(manager, pid)
        failures = []
        if result.get('status') != 'completed':
            failures.append(_goal_failure('PROVISIONAL_ARTIFACT', ['C10'], '未確認・未配分・未照合が残っています', calc_path))
        if result.get('version') != mission['plan_version'] or result.get('requirements_hash') != requirements_hash(mission):
            failures.append(_goal_failure('HASH_DRIFT', ['C07'], '計算結果が現在の目標・計画版と一致しません', calc_path))
        if envelope and envelope.get('mode') == 'vehicle-auto-v1':
            proof = read_json(resolve(manager, pid, 'vehicle_profit/automatic_proof.json'))
            if proof.get('input_hash') != digest(envelope):
                failures.append(_goal_failure('HASH_DRIFT', ['C12'], '自動抽出証跡と入力が一致しません', 'vehicle_profit/automatic_proof.json'))
            if envelope.get('recipe_learning_success') or proof.get('recipe_learning_success'):
                recon = (result.get('checks') or {}).get('source_reconciliation') or envelope.get('recipe_reconciliation') or {}
                if recon.get('passed') is not True:
                    failures.append(_goal_failure('RECONCILIATION_FAILED', ['C09'], 'レシピ参照だけではRAG成功にできません', calc_path))
        if not envelope or result.get('input_hash') != digest(envelope):
            failures.append(_goal_failure('HASH_DRIFT', ['C12'], '計算後に入力が変更されています', INPUT_PATH))
        if result.get('source_hash') != digest(sources(manager, pid)):
            failures.append(_goal_failure('HASH_DRIFT', ['C12'], '計算後に登録原本が変更されています', calc_path))
        workbook_path = c['outputs'][1]['path']
        workbook = resolve(manager, pid, workbook_path, True)
        if digest(workbook.read_bytes()) != result.get('workbook_hash'):
            failures.append(_goal_failure('HASH_DRIFT', ['C11'], '計算後にExcelが変更されています', workbook_path))
        required = ['recalculation', 'independent_profit_match', 'vehicle_summary_match', 'input_allocation_reconciled', 'coverage', 'source_controls']
        checks = result.get('checks', {})
        if not all(checks.get(k) is True for k in required):
            failures.append(_goal_failure('CHECK_MISSING', ['C07', 'C11'], '必須の計算検証が揃っていません', calc_path))
        recon = checks.get('source_reconciliation') or {}
        if recon.get('passed') is not True:
            failures.append(_goal_failure('RECONCILIATION_FAILED', ['C09'], '原本照合が全件matchedではありません', calc_path))
        if envelope:
            data = normalize_input(envelope['data'], c.get('months', []))
            failures.extend(validate_evidence_records(data, sources(manager, pid), envelope.get('confirmed', False)))
            recalculated = calculate(data)
            if recalculated['warnings']:
                failures.append(_goal_failure('ALLOCATION_UNRESOLVED', ['C10'], '再検証で欠測・推計・未配分を検出しました', calc_path))
            from openpyxl import load_workbook
            book = load_workbook(workbook, data_only=True, read_only=True)
            try:
                for idx, row in enumerate(recalculated['rows'], 2):
                    if book['サマリー'].cell(idx, 11).value != row['profit']:
                        failures.append(_goal_failure('CHECK_MISSING', ['C07'], 'Excel利益と独立再計算値が一致しません', workbook_path))
                        break
            finally:
                book.close()
        unique = []
        seen = set()
        for item in failures:
            if item['message'] in seen:
                continue
            seen.add(item['message'])
            unique.append(item)
        return unique
    except (ValueError, KeyError, TypeError, OSError) as exc:
        return [_goal_failure('CHECK_MISSING', ['C07'], '計算証拠の再検証に失敗: ' + str(exc), calc_path)]


def goal_failures(manager, pid):
    """Compatibility wrapper. Same messages and order as goal_failure_records()."""
    return [item['message'] for item in goal_failure_records(manager, pid)]


def verify_task_outputs(task, resolver):
    c=contract_of(task);failures=[]
    for output in c['outputs']:
        try:
            path=resolver(output['path'])
            if not path.is_file() or not path.stat().st_size:raise ValueError('成果物が空です')
            if path.suffix=='.json':
                body=read_json(path)
                if body.get('version')!=c['workflow_version']:raise ValueError('成果物の計画版が不一致です')
                if c['execution_kind']!='vehicle_sources' and body.get('status')!='completed':raise ValueError('検証未完了です')
        except (OSError,ValueError) as exc:failures.append(output['path']+': '+str(exc))
    return failures


def evaluate_ocr_resolution_gate(sources_list, *, project_id, store=None, feature_enabled=None):
    """SC00 と vehicle_extract の間の OCR 解決ゲート。flag 無効時は no-op。"""
    from app.ocr_invoice_adapter import evaluate_ocr_resolution_gate as _gate
    return _gate(sources_list, project_id=project_id, store=store, feature_enabled=feature_enabled)


async def execute(manager,pid,task):
    from app.upgrade_runtime import ReviewRequired
    # 採用済みアダプターの適用は vehicle_auto.prepare の明示分岐のみ。動的 exec は禁止。
    # 失敗の TRIZ 保存は project_manager._execute_task が on_vehicle_failure で行う（副作用は保存のみ）。
    mission=manager.memory.get_mission(pid);c=contract_of(task);kind=c['execution_kind']
    if c['workflow_version']!=mission['plan_version'] or c['requirements_hash']!=requirements_hash(mission):
        raise ReviewRequired('目標・計画が変更されています。計画を再生成してください')
    if kind=='vehicle_sources':return await asyncio.to_thread(prepare_sources,manager,pid,task,mission)
    if kind=='vehicle_extract':
        from app.capability_registry import ocr_feature_enabled
        if ocr_feature_enabled():
            from app.ocr_invoice_adapter import evaluate_ocr_resolution_gate
            from app.ocr_store import OcrStore
            gate = evaluate_ocr_resolution_gate(
                sources(manager, pid), project_id=pid,
                store=OcrStore(manager.memory.path), feature_enabled=True,
            )
            if gate.blocked:
                raise ReviewRequired('OCR未解決: ' + (gate.reason or '後続計算へ渡しません'))
        from app.vehicle_auto import prepare
        old=load_input(manager,pid)
        from app.vehicle_auto import current_ocr_adoption_signature
        if old and old.get('source_hash')==digest(sources(manager,pid)) and (old.get('mode')!='vehicle-auto-v1' or old.get('requirements_hash')==requirements_hash(mission) and old.get('extractor_revision')==__import__('app.vehicle_auto',fromlist=['REVISION']).REVISION and old.get('ocr_adoption_hash', '') == current_ocr_adoption_signature(manager, pid)):
            envelope=old
        else:
            try:envelope=await asyncio.to_thread(prepare,manager,pid,mission)
            except ValueError as exc:raise ReviewRequired(str(exc)) from exc
        write_json(resolve(manager,pid,c['outputs'][0]['path']),{'version':mission['plan_version'],'status':'completed','input_hash':digest(envelope),'records':len(envelope['data']['records']),'meaning':'抽出処理の完了。業務上の検査合格とは別'})
        return '原本から明細を準備しました。未解決項目を保持して暫定Excelの計算へ進みます。'
    if kind=='vehicle_verify':
        issues=goal_failures(manager,pid)
        write_json(resolve(manager,pid,c['outputs'][0]['path']),{'version':mission['plan_version'],'status':'needs_review' if issues else 'completed','issues':issues,'requirements_hash':requirements_hash(mission)})
        if issues:raise ReviewRequired('最終照合: '+' / '.join(issues[:8]))
        return '目標達成: 対象月・車両、入力原本、配分・損益、Excel再計算の照合に合格しました。'
    envelope=load_input(manager,pid)
    from app.vehicle_auto import current_ocr_adoption_signature
    if not envelope or (envelope.get('mode')=='vehicle-auto-v1' and
            (envelope.get('extractor_revision')!=__import__('app.vehicle_auto',fromlist=['REVISION']).REVISION or envelope.get('source_hash')!=digest(sources(manager,pid)) or envelope.get('requirements_hash')!=requirements_hash(mission) or envelope.get('ocr_adoption_hash', '') != current_ocr_adoption_signature(manager, pid))):
        from app.vehicle_auto import prepare
        try:envelope=await asyncio.to_thread(prepare,manager,pid,mission)
        except ValueError as exc:raise ReviewRequired(str(exc)) from exc
    snapshot=sources(manager,pid)
    if envelope.get('mode')=='vehicle-auto-v1':
        proof=read_json(resolve(manager,pid,'vehicle_profit/automatic_proof.json'))
        if proof.get('input_hash')!=digest(envelope):raise ReviewRequired('自動抽出後に入力が変更されています。再抽出してください')
    if envelope.get('source_hash')!=digest(snapshot):raise ReviewRequired('登録原本が入力確認後に変更されています。原本を再確認して入力を登録してください')
    try:
        data=normalize_input(envelope['data'],c.get('months',[]))
        issues=validate_evidence(data,snapshot,envelope.get('confirmed',False))
    except (ValueError,TypeError,KeyError) as exc:raise ReviewRequired('入力条件: '+str(exc)) from exc
    target=resolve(manager,pid,c['outputs'][1]['path']);target.parent.mkdir(parents=True,exist_ok=True)
    input_hash=digest(envelope);run_dir=target.parent/'runs'/input_hash[:16]
    recon_issues,recon_summary,recon_rows=source_reconciliation_report(data,snapshot)
    recon_path=run_dir/'source_reconciliation.json'
    recon_summary['artifact']=str(recon_path)
    write_json(recon_path,{'summary':recon_summary,'issues':recon_issues,'rows':recon_rows})
    result=await asyncio.to_thread(build_workbook,data,run_dir,source_reconciliation=recon_rows)
    if result.get('file'):target.write_bytes(Path(result['file']).read_bytes())
    if digest(load_input(manager,pid))!=input_hash or digest(sources(manager,pid))!=digest(snapshot):
        issues.append('処理中に入力または原本が変わりました')
    issues.extend(result.get('warnings',[]))
    checks=result.get('checks',{});checks['coverage']=not any(r['profit'] is None for r in calculate(data)['rows'])
    checks['source_reconciliation']=recon_summary
    checks['source_controls']=not validate_evidence(data,snapshot,envelope.get('confirmed',False))
    if recon_summary.get('passed') is not True:
        issues.extend(recon_issues)
        checks['source_controls']=False
    unallocated=any(not r.get('allocations') for r in data.get('records',[]) if not _is_policy_or_gap_record(r))
    if unallocated:
        issues.append('未配賦額が残っているため車両別確定はできません')
    completed=not issues and result.get('status')=='completed' and recon_summary.get('passed') is True and checks.get('source_controls') is True
    result.update(version=mission['plan_version'],requirements_hash=requirements_hash(mission),
                  input_hash=input_hash,source_hash=digest(snapshot),workbook_hash=digest(target.read_bytes()) if target.exists() else '',
                  status='completed' if completed else 'needs_review',issues=list(dict.fromkeys(issues)),checks=checks)
    result.pop('file',None);write_json(resolve(manager,pid,c['outputs'][0]['path']),result)
    manager.memory.add_event(pid,'vehicle_calculation','車両別・月別の計算とExcel生成を実行しました',task['id'],json.dumps({'version':mission['plan_version'],'input_hash':input_hash,'status':result['status'],'issues':result['issues']},ensure_ascii=False))
    if result['status']!='completed':raise ReviewRequired('途中版Excelを保存しました。目標未達: '+' / '.join(result['issues'][:8]))
    return '車両別・月別の損益とExcel生成・再計算が完了しました。最終照合へ進みます。'
